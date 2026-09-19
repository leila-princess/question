import re
import io
import base64
from typing import Any, Dict, List, Tuple, Optional
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree
import logging

logger = logging.getLogger(__name__)

def _docx_extract_ordered_blocks(doc: Document) -> List[Tuple[str, Any]]:
    blocks: List[Tuple[str, Any]] = []
    try:
        xml_root = etree.fromstring(doc.part.blob)
        nsmap = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        body = xml_root.find('w:body', namespaces=nsmap)
        for child in list(body):
            tag = etree.QName(child.tag).localname
            if tag == 'p':
                blocks.append(("p", child))
            elif tag == 'tbl':
                blocks.append(("tbl", child))
    except Exception:
        for p in doc.paragraphs:
            blocks.append(("p", p))
    return blocks

def _tbl_to_rows(tbl_node) -> List[List[str]]:
    nsmap = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    rows: List[List[str]] = []
    for tr in tbl_node.findall('.//w:tr', namespaces=nsmap):
        cells = []
        for tc in tr.findall('.//w:tc', namespaces=nsmap):
            texts = tc.xpath('.//w:t/text()', namespaces=nsmap)
            cells.append("".join(texts).strip())
        rows.append(cells)
    return rows

def is_heading_by_style(p_style_name: Optional[str]) -> bool:
    """通过样式判断是否为标题"""
    if not p_style_name:
        return False
    name = p_style_name.lower()
    return name.startswith("heading") or ("标题" in p_style_name)

def is_heading_by_rule(text: str) -> bool:
    """通过规则判断是否为标题"""
    text = text.strip()
    if not text:
        return False
    if len(text) > 120:
        return False
    pattern = r"^((第[一二三四五六七八九十百千]+[章节部])|([0-9]{1,2}(\.[0-9]{1,2}){0,3})|([一二三四五六七八九十]+、)).{0,60}$"
    return re.match(pattern, text) is not None

def extract_document_structure(file_bytes: bytes) -> Dict[str, Any]:
    """提取文档结构（传统方法）"""
    doc = Document(io.BytesIO(file_bytes))

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    title = paragraphs[0] if paragraphs else "未命名文档"

    blocks = _docx_extract_ordered_blocks(doc)

    headings: List[str] = []
    sections: List[Dict[str, Any]] = []
    cur_section: Optional[Dict[str, Any]] = None

    for kind, node in blocks:
        if kind == "p":
            if isinstance(node, etree._Element):
                nsmap = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                texts = node.xpath('.//w:t/text()', namespaces=nsmap)
                text = "".join(texts).strip()
                ppr = node.find('w:pPr', namespaces=nsmap)
                style = None
                if ppr is not None:
                    pstyle = ppr.find('w:pStyle', namespaces=nsmap)
                    if pstyle is not None:
                        style = pstyle.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
                is_head = is_heading_by_style(style) or is_heading_by_rule(text)
            else:
                text = getattr(node, 'text', '').strip()
                style = getattr(getattr(node, 'style', None), 'name', None)
                is_head = is_heading_by_style(style) or is_heading_by_rule(text)

            if not text:
                continue

            if is_head:
                if cur_section:
                    sections.append(cur_section)
                cur_section = {"title": text, "content": [], "tables": []}
                headings.append(text)
            else:
                if not cur_section:
                    cur_section = {"title": "前言/摘要", "content": [], "tables": []}
                    headings.append(cur_section["title"]) if not headings else None
                cur_section["content"].append(text)

        elif kind == "tbl":
            try:
                rows = _tbl_to_rows(node)
            except Exception:
                rows = []
            light_rows = [r[:6] for r in rows[:8]] if rows else []
            if not cur_section:
                cur_section = {"title": "表格", "content": [], "tables": []}
                headings.append(cur_section["title"]) if not headings else None
            cur_section["tables"].append({"rows": light_rows})

    if cur_section:
        sections.append(cur_section)

    return {
        "title": title,
        "headings": headings,
        "sections": sections,
    }

def extract_images_from_docx_bytes(docx_bytes: bytes) -> List[Dict[str, Any]]:
    """从docx中提取图片"""
    imgs = []
    try:
        z = ZipFile(io.BytesIO(docx_bytes))
        for name in z.namelist():
            if name.startswith("word/media/"):
                content = z.read(name)
                ext = os.path.splitext(name)[1].lower()
                if ext in [".jpeg", ".jpg"]:
                    ctype = "image/jpeg"
                elif ext == ".png":
                    ctype = "image/png"
                elif ext == ".gif":
                    ctype = "image/gif"
                else:
                    ctype = "application/octet-stream"
                b64 = base64.b64encode(content).decode("utf-8")
                imgs.append({"filename": os.path.basename(name), "contentType": ctype, "base64": b64})
    except Exception as e:
        logger.exception("extract_images_from_docx_bytes error: %s", e)
    return imgs

def extract_tables_and_titles(doc: Document) -> List[Dict[str, Any]]:
    """提取表格并推断标题"""
    tables_out = []
    try:
        blob = doc.part.blob
        xml_root = etree.fromstring(blob)
        nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        body = xml_root.find('w:body', namespaces=nsmap)

        children = list(body)
        for idx, child in enumerate(children):
            tag = etree.QName(child.tag).localname
            if tag == 'tbl':
                title = None
                for j in range(idx - 1, max(-1, idx - 6), -1):
                    prev = children[j]
                    if etree.QName(prev.tag).localname == 'p':
                        texts = prev.xpath('.//w:t/text()', namespaces=nsmap)
                        text = "".join(texts).strip()
                        if text and len(text) <= 120:
                            title = text
                            break
                rows = []
                for tr in child.findall('.//w:tr', namespaces=nsmap):
                    cells = []
                    for tc in tr.findall('.//w:tc', namespaces=nsmap):
                        texts = tc.xpath('.//w:t/text()', namespaces=nsmap)
                        cell_text = "".join(texts).strip()
                        cells.append(cell_text)
                    rows.append(cells)
                tables_out.append({"title": title, "rows": rows})
    except Exception as e:
        logger.exception("XML table parse fallback to python-docx: %s", e)
        for t in doc.tables:
            rows = []
            for r in t.rows:
                cells = [c.text.strip() for c in r.cells]
                rows.append(cells)
            tables_out.append({"title": None, "rows": rows})
    return tables_out

def save_word_file(final_report, metadata, timestamp, output_dir):
    """保存Word文件"""
    import os
    filename = f"{metadata.replace(' ', '_')}_{timestamp}.docx"
    filepath = os.path.join(output_dir, filename)
    doc = Document()
    title = doc.add_heading(metadata, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"标准编号: {metadata}")
    doc.add_paragraph(f"生成时间: {timestamp}")
    doc.add_paragraph()
    for paragraph in final_report.split('\n'):
        if paragraph.strip():
            p = doc.add_paragraph(paragraph)
            p.paragraph_format.space_after = Pt(6)
    doc.save(filepath)
    return filepath