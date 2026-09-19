from docx import Document

def parse_docx(file):
    """
    解析 Word 文档，提取段落和表格
    :param file: werkzeug.datastructures.FileStorage
    :return: dict
    """
    doc = Document(file)
    content = {"paragraphs": [], "tables": []}

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content["paragraphs"].append(text)

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)
        content["tables"].append(rows)

    return content
