"""
分类管理(FLGL) API — 文件上传

节点类型:
  EquipmentType, Equipment, Environment, TestProject, Indicator, TestMethod, File
"""
from flask import Blueprint, request, jsonify, current_app
from datetime import datetime
import uuid
import logging
# ==================== 文件上传接口（测试用）====================
# 在文件顶部添加必要的导入（如果没有的话）
import os
import time
from werkzeug.utils import secure_filename
#===================== 导入149文件解析文件 ===================
from docx import Document
import json
import re
from .test_extrac_output import get_table_caption, extract_entities_from_table, export_to_json, merge_entities, print_entities

logger = logging.getLogger(__name__)

flgl_upload_bp = Blueprint('flgl_upload', __name__, url_prefix='/api/FLGL')

# 配置上传文件夹（放在文件顶部，其他路由之后）
UPLOAD_FOLDER = 'uploads/flgl'
ALLOWED_EXTENSIONS = {'doc', 'docx'}

# 确保上传目录存在
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

def allowed_file(filename):
    """检查文件扩展名是否允许"""
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def process_word_document_with_output(file_path, output_dir):
    """处理Word文档，提取所有表格的实体，返回实体列表"""
    try:
        doc = Document(file_path)
    except Exception as e:
        print(f"无法打开文档: {e}")
        return []
    
    if not doc.tables:
        print("文档中没有找到表格")
        return []
    
    all_entities = []
    
    for table_idx, table in enumerate(doc.tables, 1):
        print("\n" + "=" * 80)
        
        # 获取表题注和一级设备类别
        level1_category, full_caption = get_table_caption(table)
        print(f"📋 表题注: {full_caption}")
        print(f"🏷️  一级设备类别: {level1_category}")
        
        # 提取实体
        entities = extract_entities_from_table(table, level1_category)
        all_entities.append(entities)
        
        # 打印实体
        print_entities(entities)
        
        # 导出为JSON（使用指定的输出目录）
        json_file = os.path.join(output_dir, f"entities_table_{table_idx}_{level1_category}.json")
        export_to_json(entities, json_file)
    
    print("\n" + "=" * 80)
    print("✅ 所有表格处理完成！")
    return all_entities


@flgl_upload_bp.route('/upload', methods=['POST'])
def upload_word_file():
    """
    上传 Word 文件并解析，生成图数据库节点和关系的JSON文件
    """
    try:
        logger.info("=" * 50)
        logger.info("收到文件上传请求")
        logger.info(f"请求方法: {request.method}")
        logger.info(f"请求头 Content-Type: {request.headers.get('Content-Type')}")
        
        # 检查是否有文件
        if 'file' not in request.files:
            logger.error("请求中没有 'file' 字段")
            return jsonify({
                'success': False,
                'message': '没有文件，请确保表单字段名为 "file"'
            }), 400
        
        file = request.files['file']
        logger.info(f"收到文件: {file.filename}")
        logger.info(f"文件类型: {file.content_type}")
        
        # 检查文件名
        if file.filename == '':
            logger.error("文件名为空")
            return jsonify({
                'success': False,
                'message': '文件名为空'
            }), 400
        
        # 检查文件类型
        if not allowed_file(file.filename):
            logger.error(f"不支持的文件类型: {file.filename}")
            return jsonify({
                'success': False,
                'message': f'只支持 Word 文件（.doc 或 .docx），当前文件: {file.filename}'
            }), 400
        
        # 保存文件
        filename = secure_filename(file.filename)
        timestamp = int(time.time())
        saved_filename = f"{timestamp}_{filename}"
        filepath = os.path.join(UPLOAD_FOLDER, saved_filename)
        file.save(filepath)
        
        logger.info(f"文件保存成功: {filepath}")
        logger.info(f"文件大小: {os.path.getsize(filepath)} 字节")
        
        # 创建解析结果输出目录
        output_dir = os.path.join(UPLOAD_FOLDER, f"parsed_{timestamp}")
        os.makedirs(output_dir, exist_ok=True)
        logger.info(f"创建输出目录: {output_dir}")
        
        # 解析Word文档并生成JSON
        try:
            # 调用解析函数
            all_entities = process_word_document_with_output(filepath, output_dir)
            
            # 生成合并的JSON文件
            if all_entities:
                combined_entities = merge_entities(all_entities)
                combined_json_path = os.path.join(output_dir, "all_entities_combined.json")
                export_to_json(combined_entities, combined_json_path)
                
                # 生成摘要JSON
                summary = {
                    'tables': all_entities,
                    'summary': {
                        'total_level1_categories': len(set(e['level1_category'] for e in all_entities)),
                        'total_devices': sum(len(e['devices']) for e in all_entities),
                        'total_test_items': sum(len(e['test_items']) for e in all_entities),
                        'total_indicators': sum(len(e['indicators']) for e in all_entities)
                    }
                }
                summary_json_path = os.path.join(output_dir, "all_entities_summary.json")
                with open(summary_json_path, 'w', encoding='utf-8') as f:
                    json.dump(summary, f, ensure_ascii=False, indent=2)
                
                logger.info(f"解析完成，共处理 {len(all_entities)} 个表格")
            
        except Exception as parse_error:
            logger.error(f"文件解析失败: {str(parse_error)}")
            import traceback
            traceback.print_exc()
            return jsonify({
                'success': False,
                'message': f'文件上传成功但解析失败: {str(parse_error)}'
            }), 500
        
        logger.info("=" * 50)
        
        # 返回成功响应
        return jsonify({
            'success': True,
            'message': f'文件上传并解析成功',
            'data': {
                'original_filename': file.filename,
                'saved_filename': saved_filename,
                'filepath': filepath,
                'output_dir': output_dir,
                'size': os.path.getsize(filepath),
                'parsed_tables': len(all_entities) if 'all_entities' in locals() else 0
            }
        })
        
    except Exception as e:
        logger.error(f"上传失败: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'message': f'上传失败: {str(e)}'
        }), 500