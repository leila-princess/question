"""
PDF / Word 文档标题自动识别工具（与 WJGL 文件管理共用算法）。

策略：
- Word：Heading 样式 → 编号段落规则
- PDF：书签 → PyMuPDF 规则 → OCR 回退（mode=full）

用法：
  python -m blueprints.ZSGL.WJGL1.document_heading_extract path/to/file.pdf
  from blueprints.ZSGL.WJGL1.document_heading_extract import extract_document_headings
"""
from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import tempfile
from typing import Any, Dict, List, Optional, Tuple

import docx
import fitz

logger = logging.getLogger(__name__)

_GB_TOC_KEYWORDS = (
    "范围", "引用", "术语", "定义", "试验", "设备", "过程", "结果", "报告",
    "前言", "附录", "摘要", "概述", "条件", "环境", "要求", "评定", "中断",
    "贮存", "工作", "受试", "仪器", "恢复",
)

SECTION_RE = re.compile(r"^(\d+(?:\.\d+)*|\.\d+)\s+(.+)$")


def is_text_garbled(text: str) -> bool:
    """检测文本是否疑似乱码（含 PDF 符号字体 CID 误映射）。"""
    if not text or not text.strip():
        return True
    if "\ufffd" in text or "�" in text:
        return True
    try:
        from blueprints.ZSGL.WJGL1.text_normalize import has_cid_garbled
        if has_cid_garbled(text):
            return True
    except Exception:
        pass
    total = len(text)
    if total == 0:
        return True
    replacement_count = text.count("�") + text.count("\ufffd")
    good_count = 0
    for ch in text:
        code = ord(ch)
        if ch.isalnum() or ch.isspace() or ch in "，。、《》？！：；,.!?;:()[]{}-_/\\\"'％%＋+＝=·—":
            good_count += 1
        elif 0x4E00 <= code <= 0x9FFF:
            good_count += 1
    good_ratio = good_count / total
    replacement_ratio = replacement_count / total
    return good_ratio < 0.55 or replacement_ratio > 0.02


def normalize_fullwidth(text: str) -> str:
    result = []
    for c in text:
        cp = ord(c)
        if 0xFF10 <= cp <= 0xFF19:
            result.append(chr(cp - 0xFEE0))
        elif cp == 0xFF0E:
            result.append(".")
        elif cp == 0x3000:
            result.append(" ")
        else:
            result.append(c)
    return "".join(result)


def normalize_ocr_text(text: str) -> str:
    text = normalize_fullwidth(text)
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    return text.strip()


def parse_section_number(text: str) -> Tuple[str, str]:
    m = re.match(r"^(\d+(?:\.\d+){0,4})\s*[)）.、\s]+(.+)", text)
    if m:
        return m.group(1), m.group(2).strip() or text
    m = re.match(r"^(第[一二三四五六七八九十百\d]+[章节部分篇])\s*(.*)", text)
    if m:
        return m.group(1), m.group(2).strip() or text
    return "", text


def build_toc_tree(items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    toc: List[Dict[str, Any]] = []
    stack: List[tuple] = []
    for item in items:
        level = item.get("level", 1)
        node = {
            "title": item["title"],
            "sectionNo": item.get("sectionNo", ""),
            "level": level,
            "children": [],
        }
        while stack and stack[-1][0] >= level:
            stack.pop()
        if stack:
            stack[-1][1]["children"].append(node)
        else:
            toc.append(node)
        stack.append((level, node))
    return toc


def _collect_all_titles(nodes: list) -> List[str]:
    titles: List[str] = []
    for node in nodes:
        if node.get("title"):
            titles.append(node["title"])
        if node.get("children"):
            titles.extend(_collect_all_titles(node["children"]))
    return titles


def is_toc_garbled(toc: list) -> bool:
    titles = _collect_all_titles(toc)
    if not titles:
        return True
    garbled_count = sum(1 for t in titles if is_text_garbled(t))
    garbled_ratio = garbled_count / len(titles)
    logger.info(
        "[heading_extract] 乱码检测: %s/%s (%.1f%%)",
        garbled_count, len(titles), garbled_ratio * 100,
    )
    if garbled_ratio > 0.25:
        return True
    joined = " ".join(titles)
    try:
        from blueprints.ZSGL.WJGL1.text_normalize import garbled_score, has_cid_garbled
        if has_cid_garbled(joined) or garbled_score(joined) >= 0.3:
            return True
    except Exception:
        pass
    if len(titles) >= 4:
        sample = titles[: min(15, len(titles))]
        kw_hit = sum(1 for t in sample if any(k in t for k in _GB_TOC_KEYWORDS))
        if kw_hit == 0:
            logger.info("[heading_extract] 目录无常见国标章节词，判定为 fitz 乱码")
            return True
    return False


def _validate_and_build_toc(candidates: list, skip_pages: int = 3) -> list:
    from blueprints.ZSGL.WJGL1.body_content_trim import build_toc_tree_from_candidates
    return build_toc_tree_from_candidates(candidates, skip_pages=skip_pages)


def log_toc_tree(nodes: list, prefix: str = "", indent: int = 0) -> None:
    for node in nodes:
        no = node.get("sectionNo", "")
        title = node.get("title", "")
        level = node.get("level", 1)
        children = node.get("children", [])
        logger.info(f"{prefix}{'  ' * indent}[L{level}] {no} {title}")
        if children:
            log_toc_tree(children, prefix, indent + 1)


def extract_headings_by_style(doc) -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []
    for para in doc.paragraphs:
        style_name = (para.style.name or "").lower() if para.style else ""
        text = para.text.strip()
        if not text or "heading" not in style_name:
            continue
        m = re.search(r"heading\s*(\d+)", style_name)
        level = int(m.group(1)) if m else 1
        section_no, _display = parse_section_number(text)
        items.append({"title": text, "sectionNo": section_no, "level": level})
    return items


def extract_headings_by_number(doc) -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text or len(text) > 100:
            continue
        m = re.match(r"^(\d+(?:\.\d+){0,4})\s*[)）.、\s]+(.+)", text)
        if not m:
            continue
        nums = m.group(1).split(".")
        if len(nums) > 5:
            continue
        first_num = int(nums[0])
        if first_num < 1 or first_num > 20:
            continue
        after = m.group(2).strip()
        if not after or re.match(r"^[图表附注]", after):
            continue
        items.append({
            "title": text,
            "sectionNo": m.group(1),
            "level": len(nums),
        })
    return items


def extract_toc_from_word(docx_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    从 Word 提取目录树（合并 Heading 样式 / 大纲级别 / 编号段落 / 正文检测）。
    返回 (toc_tree, method)。
    """
    try:
        from .word_native_parse import extract_word_headings_merged
        return extract_word_headings_merged(docx_path)
    except Exception as exc:
        logger.warning("[heading_extract] word_native 合并标题失败，回退旧逻辑: %s", exc)

    try:
        doc = docx.Document(docx_path)
    except Exception as exc:
        logger.warning("[heading_extract] 无法打开 Word: %s", exc)
        return [], "error"

    try:
        from_style = extract_headings_by_style(doc)
        if len(from_style) >= 3:
            logger.info("[heading_extract] Word Heading 样式: %d 条", len(from_style))
            return build_toc_tree(from_style), "heading_style"
        from_number = extract_headings_by_number(doc)
        logger.info("[heading_extract] Word 编号段落: %d 条", len(from_number))
        return build_toc_tree(from_number), "numbered_paragraph"
    except Exception as exc:
        logger.error("[heading_extract] Word 提取异常: %s", exc)
        return [], "error"


def extract_toc_from_pdf_rules(pdf_path: str) -> List[Dict[str, Any]]:
    try:
        doc = fitz.open(pdf_path)
    except Exception:
        return []

    candidates: List[Dict[str, Any]] = []
    skip_pages = 2
    for page_idx in range(len(doc)):
        page = doc[page_idx]
        blocks = page.get_text("dict")["blocks"]
        for block in blocks:
            if "lines" not in block:
                continue
            for line in block["lines"]:
                spans = line.get("spans", [])
                if not spans:
                    continue
                line_text = "".join(span.get("text", "") for span in spans)
                line_text = normalize_ocr_text(line_text)
                if not line_text:
                    continue
                m = SECTION_RE.match(line_text)
                if not m:
                    continue
                section_no = m.group(1)
                title = m.group(2).strip()
                if len(title) < 2 or len(title) > 30:
                    continue
                candidates.append({
                    "no": section_no,
                    "title": title,
                    "depth": section_no.count(".") + 1,
                    "page": page_idx + 1,
                })
    doc.close()
    return _validate_and_build_toc(candidates, skip_pages)


def extract_toc_from_pdf_ocr(pdf_path: str) -> List[Dict[str, Any]]:
    try:
        from Process_data.extractor import DocumentExtractor
        extractor = DocumentExtractor()
        result = extractor.extract(pdf_path)
        pages = result.get("pages", [])
    except Exception as exc:
        logger.warning("[heading_extract] PDF OCR 失败: %s", exc)
        return []

    candidates: List[Dict[str, Any]] = []
    for page_idx, page in enumerate(pages):
        text = normalize_ocr_text(page.get("text", ""))
        for line in text.split("\n"):
            line = line.strip()
            if not line:
                continue
            m = SECTION_RE.match(line)
            if not m:
                continue
            section_no = m.group(1)
            title = m.group(2).strip()
            if len(title) < 2 or len(title) > 30:
                continue
            candidates.append({
                "no": section_no,
                "title": title,
                "depth": section_no.count(".") + 1,
                "page": page_idx + 1,
            })
    toc = _validate_and_build_toc(candidates, skip_pages=3)
    if toc:
        logger.info("[heading_extract] PDF OCR 规则: %d 个顶级节点", len(toc))
        log_toc_tree(toc, prefix="[heading_extract/ocr] ")
    return toc


def _toc_from_pdf_bookmarks(pdf_path: str, *, simple: bool = False) -> List[Dict[str, Any]]:
    try:
        doc = fitz.open(pdf_path)
        bookmarks = doc.get_toc(simple=simple) or []
        doc.close()
    except Exception:
        return []
    if not bookmarks:
        return []

    toc: List[Dict[str, Any]] = []
    stack: List[tuple] = []
    for level, title, _page in bookmarks:
        title = (title or "").strip()
        if not title:
            continue
        m = SECTION_RE.match(title)
        if m:
            section_no = m.group(1)
            clean_title = m.group(2).strip()
        else:
            m2 = re.match(r"^(\d+(?:\.\d+)*)\s*(.*)", title)
            section_no = m2.group(1) if m2 else ""
            clean_title = (m2.group(2).strip() if m2 else title) or title
        node = {
            "title": clean_title,
            "sectionNo": section_no,
            "level": level,
            "children": [],
        }
        while stack and stack[-1][0] >= level:
            stack.pop()
        if stack:
            stack[-1][1]["children"].append(node)
        else:
            toc.append(node)
        stack.append((level, node))
    return toc


def extract_toc_from_pdf(pdf_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """完整 PDF 目录：书签 → fitz 规则 → OCR。返回 (toc, method)。"""
    bookmarks_toc = _toc_from_pdf_bookmarks(pdf_path, simple=False)
    if len(bookmarks_toc) >= 3 and not is_toc_garbled(bookmarks_toc):
        log_toc_tree(bookmarks_toc, prefix="[heading_extract/bookmarks] ")
        return bookmarks_toc, "bookmarks"

    rules_toc = extract_toc_from_pdf_rules(pdf_path)
    if rules_toc and not is_toc_garbled(rules_toc):
        log_toc_tree(rules_toc, prefix="[heading_extract/rules] ")
        return rules_toc, "fitz_rules"

    ocr_toc = extract_toc_from_pdf_ocr(pdf_path)
    if ocr_toc:
        return ocr_toc, "ocr_rules"
    if rules_toc:
        return rules_toc, "fitz_rules_fallback"
    if bookmarks_toc:
        return bookmarks_toc, "bookmarks_fallback"
    return [], "none"


def extract_toc_from_pdf_fast(pdf_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """快速 PDF 目录：书签 + fitz 规则，不触发 OCR。"""
    bookmarks_toc = _toc_from_pdf_bookmarks(pdf_path, simple=True)
    if bookmarks_toc and not is_toc_garbled(bookmarks_toc):
        log_toc_tree(bookmarks_toc, prefix="[heading_extract/fast/bookmarks] ")
        return bookmarks_toc, "bookmarks"
    rules_toc = extract_toc_from_pdf_rules(pdf_path)
    if rules_toc:
        log_toc_tree(rules_toc, prefix="[heading_extract/fast/rules] ")
    return rules_toc or [], "fitz_rules" if rules_toc else "none"


def resolve_document_path(file_path: str) -> Tuple[str, str, Optional[str]]:
    """
    解析输入路径；.doc 经 LibreOffice 转为 .docx（不转 PDF、不走 MinerU）。
    返回 (resolved_path, ext, temp_to_cleanup)。
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext != ".doc":
        return file_path, ext, None
    try:
        import tempfile
        from .word_native_parse import _ensure_docx_path

        out_dir = tempfile.mkdtemp(prefix="heading_extract_docx_")
        docx_path, converted = _ensure_docx_path(file_path, work_dir=out_dir)
        if converted:
            return docx_path, ".docx", out_dir
        return docx_path, ".docx", None
    except Exception as exc:
        logger.warning("[heading_extract] .doc 转 docx 失败: %s", exc)
        return file_path, ext, None


def flatten_toc_to_headings(
    toc: List[Dict[str, Any]],
    *,
    source: str = "wjgl",
) -> List[Dict[str, Any]]:
    """将嵌套目录树展平为前端 full_headings 格式（文档序）。"""
    out: List[Dict[str, Any]] = []
    counter = 0

    def walk(nodes: List[Dict[str, Any]]) -> None:
        nonlocal counter
        for node in nodes or []:
            if not isinstance(node, dict):
                continue
            counter += 1
            section_no = str(node.get("sectionNo") or "").strip()
            title = str(node.get("title") or "").strip()
            if not title and not section_no:
                continue
            out.append({
                "section_id": f"S{counter:03d}",
                "sectionNo": section_no,
                "title": title,
                "level": int(node.get("level") or 1),
                "start_page": node.get("page") or node.get("start_page"),
                "end_page": node.get("end_page"),
                "snippet": "",
                "source": source,
            })
            walk(list(node.get("children") or []))

    walk(toc)
    return out


def extract_headings_from_text(text: str) -> List[Dict[str, Any]]:
    """从纯文本 / Markdown 识别标题：Markdown # 标题 + 国标编号行。"""
    items: List[Dict[str, Any]] = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line or len(line) > 120:
            continue

        md = re.match(r"^(#{1,6})\s+(.+)$", line)
        if md:
            level = len(md.group(1))
            title = md.group(2).strip().strip("*")
            section_no, display = parse_section_number(title)
            items.append({
                "title": display or title,
                "sectionNo": section_no,
                "level": level,
            })
            continue

        m = re.match(r"^(\d+(?:\.\d+){0,4})\s*[)）.、\s]+(.+)", line)
        if m:
            nums = m.group(1).split(".")
            if len(nums) <= 5 and 1 <= int(nums[0]) <= 20:
                after = m.group(2).strip().strip("*")
                if after and not re.match(r"^[图表附注]", after) and len(after) <= 80:
                    items.append({
                        "title": after,
                        "sectionNo": m.group(1),
                        "level": len(nums),
                    })
            continue

        proj = re.match(r"^试验项目[一二三四五六七八九十0-9]+[：:]\s*(.+)", line)
        if proj:
            items.append({
                "title": proj.group(1).strip(),
                "sectionNo": "",
                "level": 2,
            })

    return items


def extract_toc_from_text_file(file_path: str) -> Tuple[List[Dict[str, Any]], str]:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    try:
        from .document_structure import analyze_document_structure

        structure = analyze_document_structure(full_text=text, file_path=file_path)
        tree = structure.toc_display or structure.heading_tree
        if tree:
            return tree, structure.method or "document_structure"
    except Exception as exc:
        logger.debug("[heading_extract] document_structure 文本失败: %s", exc)

    items = extract_headings_from_text(text)
    if len(items) >= 3:
        return build_toc_tree(items), "markdown_headings"
    try:
        from blueprints.FLGL.GetTestMethod.parser import split_standard_sections
        sections = split_standard_sections(text)
    except Exception:
        sections = []
    flat_items = []
    for sec in sections or []:
        heading = str(sec.get("heading") or sec.get("title") or "").strip()
        if not heading or heading in ("正文", "前言"):
            continue
        sn, tt = parse_section_number(heading)
        flat_items.append({
            "title": tt or heading,
            "sectionNo": sn,
            "level": int(sec.get("level") or 1),
        })
    if flat_items:
        return build_toc_tree(flat_items), "parser_sections"
    return build_toc_tree(items), "markdown_headings"


def extract_document_headings(
    file_path: str,
    *,
    mode: str = "full",
    filter_from_chapter_one: bool = True,
) -> Dict[str, Any]:
    """
    统一入口：自动识别 PDF / Word 标题。

    mode:
      - full: PDF 含 OCR 回退（质量优先，较慢）
      - fast: PDF 仅书签 + fitz 规则（速度优先）

    返回:
      toc, headings, heading_count, extract_method, file_ext
    """
    if not file_path or not os.path.isfile(file_path):
        raise FileNotFoundError(file_path or "empty path")

    resolved, ext, temp_dir = resolve_document_path(os.path.abspath(file_path))
    extract_method = "unknown"
    toc: List[Dict[str, Any]] = []

    try:
        if ext in (".docx", ".doc"):
            toc, extract_method = extract_toc_from_word(resolved)
        elif ext == ".pdf":
            if str(mode or "").lower() == "fast":
                toc, extract_method = extract_toc_from_pdf_fast(resolved)
            else:
                toc, extract_method = extract_toc_from_pdf(resolved)
        elif ext in (".md", ".txt"):
            toc, extract_method = extract_toc_from_text_file(resolved)
        else:
            raise ValueError(f"不支持的文件类型: {ext}")

        if filter_from_chapter_one and toc:
            from blueprints.ZSGL.WJGL1.body_content_trim import filter_toc_from_chapter_one
            toc = filter_toc_from_chapter_one(toc)

        headings = flatten_toc_to_headings(toc, source="wjgl")
        return {
            "file_path": file_path,
            "resolved_path": resolved,
            "file_ext": ext,
            "mode": mode,
            "extract_method": extract_method,
            "heading_count": len(headings),
            "toc": toc,
            "headings": headings,
            "full_headings": headings,
        }
    finally:
        if temp_dir and os.path.isdir(temp_dir):
            try:
                import shutil
                shutil.rmtree(temp_dir, ignore_errors=True)
            except Exception:
                pass


def _cli_main() -> None:
    parser = argparse.ArgumentParser(description="PDF/Word 文档标题自动识别（WJGL 算法）")
    parser.add_argument("file", help="文件路径 (.pdf / .doc / .docx / .md / .txt)")
    parser.add_argument("--mode", choices=("full", "fast"), default="full", help="PDF 识别模式")
    parser.add_argument("--no-filter", action="store_true", help="不从第 1 章开始裁剪目录")
    parser.add_argument("-o", "--output", help="输出 JSON 路径")
    parser.add_argument("--flat-only", action="store_true", help="仅输出 flat headings 列表")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    result = extract_document_headings(
        args.file,
        mode=args.mode,
        filter_from_chapter_one=not args.no_filter,
    )
    payload = result["headings"] if args.flat_only else {
        "file": result["file_path"],
        "extract_method": result["extract_method"],
        "heading_count": result["heading_count"],
        "toc": result["toc"],
        "headings": result["headings"],
    }
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"[OK] 已写入 {args.output}  headings={result['heading_count']}  method={result['extract_method']}")
    else:
        print(text)


if __name__ == "__main__":
    _backend = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    if _backend not in sys.path:
        sys.path.insert(0, _backend)
    _repo = os.path.dirname(_backend)
    if _repo not in sys.path:
        sys.path.insert(0, _repo)
    _cli_main()
