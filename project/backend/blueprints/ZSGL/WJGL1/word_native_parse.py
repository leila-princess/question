"""
Word 原生解析（.docx / .doc）：按文档顺序保留段落、标题、表格、图片，不经过 MinerU/PDF。

.doc 会先经 LibreOffice 转为 .docx 再解析（仅格式转换，不做 OCR）。
"""
from __future__ import annotations

import base64
import html
import logging
import os
import re
import tempfile
import uuid
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

try:
    import docx
    from docx.document import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    docx = None
    DocxDocument = None
    qn = None
    Table = None
    Paragraph = None

HEADING_STYLE_RE = re.compile(r"(?:heading|标题)\s*(\d+)", re.I)
NUMBERED_HEADING_RE = re.compile(r"^(\d+(?:\.\d+){0,5})\s*[)）.、\s]+(.+)$")
CHAPTER_HEADING_RE = re.compile(r"^第[一二三四五六七八九十百零0-9]+[章节部分篇]\s*(.*)$")
_PAREN_LETTER_RE = re.compile(r"^[（(]?([a-zA-Z])[)）.、\s]+(.+)$")


def _paragraph_numpr_info(para: Paragraph) -> Tuple[bool, Optional[str], Optional[str]]:
    """读取段落 w:numPr，返回 (has_numpr, num_id, ilvl)。"""
    if para is None or qn is None:
        return False, None, None
    p_pr = para._element.pPr
    if p_pr is None:
        return False, None, None
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        return False, None, None
    num_id_node = num_pr.find(qn("w:numId"))
    ilvl_node = num_pr.find(qn("w:ilvl"))
    num_id = num_id_node.get(qn("w:val")) if num_id_node is not None else None
    ilvl = ilvl_node.get(qn("w:val")) if ilvl_node is not None else "0"
    return True, num_id, ilvl


def _int_to_letter(n: int, upper: bool = False) -> str:
    n = max(1, n)
    s = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        ch = chr(ord("A" if upper else "a") + rem)
        s = ch + s
    return s


def _format_number_value(num_fmt: str, value: int) -> str:
    fmt = (num_fmt or "decimal").lower()
    if fmt in ("decimal", "decimalzero", "decimalenclosedcircle"):
        return str(value)
    if fmt in ("lowerletter",):
        return _int_to_letter(value, upper=False)
    if fmt in ("upperletter",):
        return _int_to_letter(value, upper=True)
    if fmt in ("lowerroman", "upperroman"):
        pairs = [
            (1000, "m"), (900, "cm"), (500, "d"), (400, "cd"),
            (100, "c"), (90, "xc"), (50, "l"), (40, "xl"),
            (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"),
        ]
        n = max(1, value)
        out = []
        for arabic, roman in pairs:
            while n >= arabic:
                out.append(roman)
                n -= arabic
        s = "".join(out)
        return s.upper() if fmt == "upperroman" else s
    if fmt in ("chinesecounting", "chinesecountingthousand", "japanesecounting"):
        digits = "零一二三四五六七八九十"
        if 1 <= value <= 10:
            return digits[value]
        return str(value)
    return str(value)


class _WordNumberingResolver:
    """轻量 Word numbering.xml 解析，为列表/标题段落生成可见编号前缀。"""

    def __init__(self, doc: DocxDocument):
        self._counters: Dict[Tuple[str, int], int] = {}
        self._abstract: Dict[str, Dict[int, Dict[str, str]]] = {}
        self._num_to_abs: Dict[str, str] = {}
        self._load(doc)

    def _load(self, doc: DocxDocument) -> None:
        if doc is None or qn is None:
            return
        try:
            numbering_part = doc.part.numbering_part
            root = numbering_part.element
        except Exception:
            return
        for abs_num in root.findall(qn("w:abstractNum")):
            abs_id = abs_num.get(qn("w:abstractNumId"))
            if abs_id is None:
                continue
            levels: Dict[int, Dict[str, str]] = {}
            for lvl in abs_num.findall(qn("w:lvl")):
                ilvl_s = lvl.get(qn("w:ilvl")) or "0"
                try:
                    ilvl_i = int(ilvl_s)
                except ValueError:
                    ilvl_i = 0
                num_fmt_el = lvl.find(qn("w:numFmt"))
                lvl_text_el = lvl.find(qn("w:lvlText"))
                levels[ilvl_i] = {
                    "numFmt": (num_fmt_el.get(qn("w:val")) if num_fmt_el is not None else "decimal") or "decimal",
                    "lvlText": (lvl_text_el.get(qn("w:val")) if lvl_text_el is not None else "%1.") or "%1.",
                }
            self._abstract[abs_id] = levels
        for num in root.findall(qn("w:num")):
            num_id = num.get(qn("w:numId"))
            abs_ref = num.find(qn("w:abstractNumId"))
            if num_id and abs_ref is not None:
                abs_id = abs_ref.get(qn("w:val"))
                if abs_id:
                    self._num_to_abs[num_id] = abs_id

    def prefix_for_paragraph(self, para: Paragraph) -> str:
        has_num, num_id, ilvl_s = _paragraph_numpr_info(para)
        if not has_num or not num_id or num_id == "0":
            return ""
        try:
            ilvl = int(ilvl_s or "0")
        except ValueError:
            ilvl = 0
        abs_id = self._num_to_abs.get(str(num_id))
        if not abs_id:
            return ""
        levels = self._abstract.get(abs_id) or {}
        lvl_def = levels.get(ilvl) or levels.get(0) or {}
        num_fmt = str(lvl_def.get("numFmt") or "decimal")
        if num_fmt == "bullet":
            return ""
        key = (str(num_id), ilvl)
        self._counters[key] = self._counters.get(key, 0) + 1
        for deeper in range(ilvl + 1, 9):
            self._counters.pop((str(num_id), deeper), None)
        value = self._counters[key]
        lvl_text = str(lvl_def.get("lvlText") or "%1.")
        rendered = lvl_text
        for i in range(ilvl, -1, -1):
            i_def = levels.get(i) or levels.get(0) or {}
            i_fmt = str(i_def.get("numFmt") or "decimal")
            i_key = (str(num_id), i)
            i_val = self._counters.get(i_key, 1)
            token = f"%{i + 1}"
            rendered = rendered.replace(token, _format_number_value(i_fmt, i_val))
        rendered = rendered.replace("%1", _format_number_value(num_fmt, value))
        return rendered.strip()


def _display_text_for_paragraph(para: Paragraph, numbering: _WordNumberingResolver) -> str:
    raw = (para.text or "").strip()
    if not raw:
        return ""
    prefix = numbering.prefix_for_paragraph(para)
    if not prefix:
        return raw
    if raw.startswith(prefix):
        return raw
    sep = "" if prefix.endswith((".", ")", "）", "、")) else " "
    return f"{prefix}{sep}{raw}".strip()


def iter_word_body_enriched(doc: DocxDocument) -> List[Dict[str, Any]]:
    """单次遍历：生成带编号可见文本与 char 偏移的段落映射。"""
    numbering = _WordNumberingResolver(doc)
    paragraph_map: List[Dict[str, Any]] = []
    char_pos = 0
    para_index = 0

    for block in _iter_body_blocks(doc):
        if isinstance(block, Paragraph):
            display = _display_text_for_paragraph(block, numbering)
            if not display:
                imgs = _image_parts_from_paragraph(block, doc)
                if imgs:
                    para_index += 1
                continue
            level = detect_paragraph_heading_level(block, display)
            section_no = ""
            m = NUMBERED_HEADING_RE.match(display)
            if m:
                section_no = m.group(1)
            char_start = char_pos
            char_end = char_start + len(display)
            has_num, num_id, ilvl = _paragraph_numpr_info(block)
            paragraph_map.append({
                "para_index": para_index,
                "char_start": char_start,
                "char_end": char_end,
                "text": display,
                "raw_text": (block.text or "").strip(),
                "is_heading": bool(level),
                "level": int(level or 0),
                "section_no": section_no,
                "has_numpr": has_num,
                "num_id": num_id,
                "ilvl": ilvl,
            })
            char_pos = char_end + 2
            para_index += 1
        elif isinstance(block, Table):
            rows_text = []
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    rows_text.append(" | ".join(cells))
            if rows_text:
                tbl_text = "【表格】\n" + "\n".join(rows_text)
                char_start = char_pos
                char_end = char_start + len(tbl_text)
                paragraph_map.append({
                    "para_index": para_index,
                    "char_start": char_start,
                    "char_end": char_end,
                    "text": tbl_text,
                    "raw_text": tbl_text,
                    "is_heading": False,
                    "level": 0,
                    "section_no": "",
                    "is_table": True,
                })
                char_pos = char_end + 2
                para_index += 1
    return paragraph_map


def build_word_heading_candidates_from_map(
    paragraph_map: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """从 paragraph_map 提取标题候选（offset 与 full_text 同源）。"""
    out: List[Dict[str, Any]] = []
    for entry in paragraph_map or []:
        if not entry.get("is_heading"):
            continue
        title = str(entry.get("text") or "").strip()
        if not title:
            continue
        out.append({
            "offset": int(entry.get("char_start") or 0),
            "level": max(1, int(entry.get("level") or 1)),
            "section_no": str(entry.get("section_no") or ""),
            "title": title,
            "heading": title,
            "source": "word_paragraph_map",
            "confidence": 0.92,
            "para_index": entry.get("para_index"),
        })
    return out


def _ensure_docx_path(file_path: str, work_dir: Optional[str] = None) -> Tuple[str, bool]:
    """返回可用于 python-docx 的路径；.doc 时先转 docx。"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return os.path.abspath(file_path), False
    if ext != ".doc":
        raise ValueError(f"不支持的 Word 类型: {ext}")
    out_dir = work_dir or tempfile.mkdtemp(prefix="word_to_docx_")
    os.makedirs(out_dir, exist_ok=True)
    try:
        from utils.office_to_pdf import get_soffice_path
    except ImportError:
        from backend.utils.office_to_pdf import get_soffice_path

    soffice = get_soffice_path()
    if not soffice:
        raise RuntimeError(
            "解析 .doc 需要 LibreOffice 将文件转为 .docx（不经过 MinerU）。"
            "请安装 LibreOffice 并配置 LIBREOFFICE_PATH。"
        )
    import subprocess

    cmd = [
        soffice,
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        "--convert-to",
        "docx",
        "--outdir",
        out_dir,
        os.path.abspath(file_path),
    ]
    logger.info("[word_native] .doc → .docx: %s", os.path.basename(file_path))
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180, cwd=out_dir)
    if proc.returncode != 0:
        err = (proc.stderr or proc.stdout or "").strip()
        raise RuntimeError(f"LibreOffice doc→docx 失败: {err[:400]}")
    base = os.path.splitext(os.path.basename(file_path))[0]
    docx_path = os.path.join(out_dir, f"{base}.docx")
    if not os.path.isfile(docx_path):
        raise RuntimeError("LibreOffice 未生成 .docx 文件")
    return docx_path, True


def _paragraph_outline_level(para: Paragraph) -> Optional[int]:
    p_pr = para._element.pPr
    if p_pr is None:
        return None
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is not None and outline.get(qn("w:val")) is not None:
        try:
            return int(outline.get(qn("w:val"))) + 1
        except (TypeError, ValueError):
            return None
    return None


def _paragraph_style_level(para: Paragraph) -> Optional[int]:
    style_name = (para.style.name or "").strip() if para.style else ""
    if not style_name:
        return None
    m = HEADING_STYLE_RE.search(style_name)
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            return 1
    lower = style_name.lower()
    if lower in ("title", "subtitle", "标题"):
        return 1
    return None


def _paragraph_numbered_heading_level(text: str) -> Optional[int]:
    text = (text or "").strip()
    if not text or len(text) > 120:
        return None
    m = NUMBERED_HEADING_RE.match(text)
    if m:
        nums = m.group(1).split(".")
        if 1 <= int(nums[0]) <= 30 and len(nums) <= 6:
            return len(nums)
    if CHAPTER_HEADING_RE.match(text):
        return 1
    return None


def detect_paragraph_heading_level(
    para: Paragraph,
    display_text: str = "",
) -> Optional[int]:
    for fn in (_paragraph_style_level, _paragraph_outline_level):
        lv = fn(para)
        if lv:
            return min(max(lv, 1), 6)
    text = (display_text or para.text or "").strip()
    num_lv = _paragraph_numbered_heading_level(text)
    if num_lv:
        return num_lv
    if text and len(text) <= 80:
        runs = list(para.runs)
        if runs and all(r.bold for r in runs if (r.text or "").strip()):
            if num_lv := _paragraph_numbered_heading_level(text):
                return num_lv
            if len(text) <= 40 and not text.endswith(("。", "；", "，", ".", ";", ",")):
                return 2
    return None


def _section_anchor_id(section_no: str) -> str:
    return "wjgl-sec-" + (section_no or "").lstrip(".").replace(".", "-")


def _runs_to_html(para: Paragraph) -> str:
    chunks: List[str] = []
    for run in para.runs:
        t = run.text or ""
        if not t:
            continue
        t = html.escape(t)
        if run.bold:
            t = f"<strong>{t}</strong>"
        if run.italic:
            t = f"<em>{t}</em>"
        if run.underline:
            t = f"<u>{t}</u>"
        chunks.append(t)
    if chunks:
        return "".join(chunks)
    return html.escape((para.text or "").strip())


def _table_to_html(table: Table) -> str:
    rows_html: List[str] = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cell_text = html.escape((cell.text or "").strip()).replace("\n", "<br>")
            cells.append(f"<td>{cell_text}</td>")
        if any(c.strip() for c in (cell.text for cell in row.cells)):
            rows_html.append(f"<tr>{''.join(cells)}</tr>")
    if not rows_html:
        return ""
    return f'<table class="doc-table word-native-table">{"".join(rows_html)}</table>'


def _escape_md_cell(text: str) -> str:
    return (text or "").replace("|", "\\|").replace("\n", " ").strip()


def _runs_to_markdown(para: Paragraph) -> str:
    chunks: List[str] = []
    for run in para.runs:
        t = run.text or ""
        if not t:
            continue
        t = t.replace("\\", "\\\\").replace("|", "\\|").replace("*", "\\*").replace("_", "\\_")
        if run.bold:
            t = f"**{t}**"
        if run.italic:
            t = f"*{t}*"
        chunks.append(t)
    if chunks:
        return "".join(chunks)
    return _escape_md_cell((para.text or "").strip())


def _table_to_gfm(table: Table) -> str:
    rows: List[List[str]] = []
    for row in table.rows:
        cells = [_escape_md_cell(c.text) for c in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    lines: List[str] = []
    head = rows[0] + [""] * (width - len(rows[0]))
    lines.append("| " + " | ".join(head) + " |")
    lines.append("| " + " | ".join(["---"] * width) + " |")
    for row in rows[1:]:
        padded = row + [""] * (width - len(row))
        lines.append("| " + " | ".join(padded) + " |")
    return "\n".join(lines)


def _figure_markdown(blob: bytes, content_type: str, index: int, caption: str = "") -> str:
    b64 = base64.b64encode(blob).decode("ascii")
    alt = caption or f"图{index}"
    return f"![{alt}](data:{content_type};base64,{b64})"


def _iter_body_blocks(doc: DocxDocument):
    body = doc.element.body
    for child in body:
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _paragraphs_in_table(table: Table) -> List[Paragraph]:
    """表格单元格内段落（Visio/OLE 图常嵌在单元格中）。"""
    paras: List[Paragraph] = []
    if Table is None or Paragraph is None:
        return paras
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                paras.append(para)
    return paras


_METAFILE_CONTENT_TYPES = frozenset({
    "image/emf", "image/x-emf", "image/wmf", "image/x-wmf",
    "application/x-msmetafile", "application/emf", "application/wmf",
})
_OLE_CONTENT_TYPES = frozenset({
    "application/vnd.openxmlformats-officedocument.oleObject",
    "application/vnd.ms-office.visio.drawing",
    "application/x-visio",
})
_OLE_COMPOUND_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_BROWSER_IMAGE_MIMES = frozenset({
    "image/png", "image/jpeg", "image/jpg", "image/gif", "image/webp", "image/svg+xml",
})


def _is_ole_compound(blob: bytes) -> bool:
    return bool(blob) and blob[:8] == _OLE_COMPOUND_MAGIC


def _is_non_displayable_part(content_type: str, blob: bytes) -> bool:
    ctype = (content_type or "").lower()
    if _is_ole_compound(blob):
        return True
    if any(x in ctype for x in ("oleobject", "visio.drawing", "msword", "ms-excel")):
        return True
    return ctype in _OLE_CONTENT_TYPES


def _detect_raster_mime(blob: bytes) -> str:
    if not blob:
        return ""
    if blob[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if blob[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if blob[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    if len(blob) > 12 and blob[:4] == b"RIFF" and blob[8:12] == b"WEBP":
        return "image/webp"
    head = blob[:256].lstrip()
    if head.startswith(b"<svg") or head.startswith(b"<?xml"):
        return "image/svg+xml"
    return ""


def _normalize_browser_image_mime(content_type: str, blob: bytes) -> str:
    detected = _detect_raster_mime(blob)
    if detected:
        return detected
    ctype = (content_type or "").split(";")[0].strip().lower()
    if ctype in _BROWSER_IMAGE_MIMES:
        return "image/jpeg" if ctype == "image/jpg" else ctype
    return ""


def _try_pillow_to_png(blob: bytes) -> Tuple[Optional[bytes], str]:
    try:
        from PIL import Image
        import io

        img = Image.open(io.BytesIO(blob))
        if img.mode not in ("RGB", "RGBA", "L", "P"):
            img = img.convert("RGBA")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        out = buf.getvalue()
        return (out if out else None), "image/png"
    except Exception:
        return None, ""


def _try_magick_metafile_to_png(blob: bytes, ext: str, work_dir: str) -> Tuple[Optional[bytes], str]:
    import subprocess

    try:
        try:
            from utils.office_to_pdf import get_magick_path
        except ImportError:
            from backend.utils.office_to_pdf import get_magick_path
    except ImportError:
        get_magick_path = None  # type: ignore

    magick = get_magick_path() if get_magick_path else None
    if not magick:
        return None, ""
    tag = uuid.uuid4().hex[:8]
    in_path = os.path.join(work_dir, f"word_meta_{tag}{ext}")
    out_path = os.path.join(work_dir, f"word_meta_{tag}.png")
    try:
        with open(in_path, "wb") as wf:
            wf.write(blob)
        # IMv7: magick -density 200 input.emf output.png
        cmd = [magick, "-density", "200", in_path, out_path]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=90)
        if proc.returncode != 0 or not os.path.isfile(out_path):
            logger.debug(
                "[word_native] ImageMagick metafile->png failed: %s",
                (proc.stderr or proc.stdout or "")[:300],
            )
            return None, ""
        with open(out_path, "rb") as pf:
            data = pf.read()
        if _detect_raster_mime(data) == "image/png":
            return data, "image/png"
    except Exception as exc:
        logger.debug("[word_native] ImageMagick convert skipped: %s", exc)
    return None, ""


def _part_blob_from_rel_id(doc: DocxDocument, rel_id: str) -> Tuple[Optional[bytes], str]:
    if not rel_id or doc is None:
        return None, ""
    try:
        part = doc.part.related_parts.get(rel_id)
        if part is None:
            rel = doc.part.rels.get(rel_id)
            if rel is not None:
                part = rel.target_part
        if part is None:
            return None, ""
        blob = part.blob
        ctype = str(getattr(part, "content_type", None) or "application/octet-stream")
        return (blob if blob else None), ctype
    except Exception as exc:
        logger.debug("[word_native] rel part read failed rel=%s: %s", rel_id, exc)
        return None, ""


def _guess_metafile_ext(blob: bytes, content_type: str) -> str:
    ctype = (content_type or "").lower()
    if "emf" in ctype:
        return ".emf"
    if "wmf" in ctype:
        return ".wmf"
    if blob[:4] == b"\x01\x00\x00\x00":
        return ".emf"
    if len(blob) >= 2 and blob[:2] in (b"\xd7\xcd", b"\x01\x00"):
        return ".wmf"
    return ""


def _is_metafile_content_type(content_type: str, blob: bytes) -> bool:
    ctype = (content_type or "").lower()
    if ctype in _METAFILE_CONTENT_TYPES or "emf" in ctype or "wmf" in ctype:
        return True
    return bool(_guess_metafile_ext(blob, content_type))


def _extract_metafile_from_ole(blob: bytes) -> Tuple[Optional[bytes], str]:
    """从 Visio/OLE 复合文档中提取 EMF/WMF 预览流。"""
    if not _is_ole_compound(blob):
        return None, ""
    try:
        import io

        import olefile

        ole = olefile.OleFileIO(io.BytesIO(blob))
        for stream in ole.listdir():
            try:
                data = ole.openstream(stream).read()
            except Exception:
                continue
            if not data or len(data) < 32:
                continue
            ext = _guess_metafile_ext(data, "")
            if ext:
                ctype = "image/x-emf" if ext == ".emf" else "image/x-wmf"
                return data, ctype
    except Exception as exc:
        logger.debug("[word_native] olefile extract skipped: %s", exc)

    pos = 0
    while pos < len(blob) - 88:
        idx = blob.find(b"\x01\x00\x00\x00", pos)
        if idx < 0:
            break
        chunk = blob[idx: min(len(blob), idx + 512 * 1024)]
        if _guess_metafile_ext(chunk, ""):
            return chunk, "image/x-emf"
        pos = idx + 4
    return None, ""


def _convert_metafile_to_png(blob: bytes, content_type: str, work_dir: str) -> Tuple[Optional[bytes], str]:
    """EMF/WMF（Visio/OLE 预览）→ PNG，便于 HTML/浏览器展示。"""
    primary = _guess_metafile_ext(blob, content_type)
    exts: List[str] = []
    for ext in (primary, ".emf", ".wmf"):
        if ext and ext not in exts:
            exts.append(ext)
    if not exts:
        return None, ""

    for ext in exts:
        png, ctype = _try_magick_metafile_to_png(blob, ext, work_dir)
        if png:
            return png, ctype

    try:
        try:
            from utils.office_to_pdf import get_soffice_path
        except ImportError:
            from backend.utils.office_to_pdf import get_soffice_path

        soffice = get_soffice_path()
        if not soffice:
            return None, ""
        import subprocess

        for ext in exts:
            tag = uuid.uuid4().hex[:8]
            in_path = os.path.join(work_dir, f"word_meta_{tag}{ext}")
            out_dir = os.path.join(work_dir, f"word_meta_out_{tag}")
            os.makedirs(out_dir, exist_ok=True)
            with open(in_path, "wb") as wf:
                wf.write(blob)
            cmd = [
                soffice,
                "--headless",
                "--nologo",
                "--nofirststartwizard",
                "--convert-to",
                "png",
                "--outdir",
                out_dir,
                in_path,
            ]
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=120, cwd=out_dir)
            if proc.returncode != 0:
                logger.debug(
                    "[word_native] metafile->png failed ext=%s stderr=%s",
                    ext,
                    (proc.stderr or proc.stdout or "")[:200],
                )
                continue
            base = os.path.splitext(os.path.basename(in_path))[0]
            png_path = os.path.join(out_dir, f"{base}.png")
            if os.path.isfile(png_path):
                with open(png_path, "rb") as pf:
                    data = pf.read()
                if _detect_raster_mime(data):
                    return data, "image/png"
    except Exception as exc:
        logger.debug("[word_native] metafile convert skipped: %s", exc)
    return None, ""


def _resolve_figure_blob(
    blob: bytes,
    content_type: str,
    *,
    work_dir: str = "",
) -> Tuple[Optional[bytes], str]:
    """解析为可显示图片；OLE 先尝试提取内部 EMF/WMF。"""
    if not blob:
        return None, ""
    if _is_ole_compound(blob) or _is_non_displayable_part(content_type, blob):
        inner, inner_ctype = _extract_metafile_from_ole(blob)
        if inner:
            blob, content_type = inner, inner_ctype
        elif _is_ole_compound(blob):
            return None, ""

    browser_mime = _normalize_browser_image_mime(content_type, blob)
    if browser_mime:
        return blob, browser_mime

    if _is_metafile_content_type(content_type, blob):
        if work_dir:
            png, png_ctype = _convert_metafile_to_png(blob, content_type, work_dir)
            if png and _detect_raster_mime(png):
                return png, png_ctype
        return None, ""

    png, png_ctype = _try_pillow_to_png(blob)
    if png and _detect_raster_mime(png):
        return png, png_ctype

    return None, ""


def _normalize_figure_blob(
    blob: bytes,
    content_type: str,
    *,
    work_dir: str = "",
) -> Tuple[Optional[bytes], str]:
    """仅返回浏览器可显示的 PNG/JPEG 等；EMF/OLE 必须转换成功，否则丢弃。"""
    norm, mime = _resolve_figure_blob(blob, content_type, work_dir=work_dir)
    if not norm and _is_metafile_content_type(content_type, blob):
        logger.info("[word_native] Visio/EMF 预览未能转为 PNG（请安装 LibreOffice 或 ImageMagick）")
    return norm, mime


def _visio_caption_from_paragraph(para: Paragraph) -> str:
    if para is None or qn is None:
        return ""
    try:
        for ole in para._element.xpath(".//*[local-name()='OLEObject']"):
            prog = str(ole.get("ProgID") or ole.get("Type") or "")
            if "visio" in prog.lower():
                return "Visio图"
    except Exception:
        pass
    return ""


def _image_rel_ids_from_paragraph(para: Paragraph) -> List[str]:
    """
    按文档顺序收集段落内图片关系 Id：
    - DrawingML a:blip（r:embed / r:link）
    - VML v:imagedata（Visio/OLE 预览 EMF/WMF）
    """
    if para is None or qn is None:
        return []
    rel_ids: List[str] = []
    seen: set = set()
    el = para._element

    def _add(rid: Optional[str]) -> None:
        rid = str(rid or "").strip()
        if rid and rid not in seen:
            seen.add(rid)
            rel_ids.append(rid)

    for blip in el.xpath(".//*[local-name()='blip']"):
        _add(blip.get(qn("r:embed")))
        _add(blip.get(qn("r:link")))

    for imagedata in el.xpath(".//*[local-name()='imagedata']"):
        _add(imagedata.get(qn("r:id")))
        _add(imagedata.get(qn("r:embed")))

    # legacy w:pict / binData
    for pict in el.xpath(".//*[local-name()='pict']"):
        for imagedata in pict.xpath(".//*[local-name()='imagedata']"):
            _add(imagedata.get(qn("r:id")))
            _add(imagedata.get(qn("r:embed")))

    return rel_ids


def _image_parts_from_paragraph(
    para: Paragraph,
    doc: DocxDocument,
    *,
    work_dir: str = "",
) -> List[Tuple[bytes, str, str]]:
    """
    提取段落内嵌图片（含 Visio/OLE 的 EMF/WMF 预览图）。
    返回 [(blob, content_type, caption_hint), ...]（仅转换成功的）
    """
    out: List[Tuple[bytes, str, str]] = []
    if docx is None or doc is None:
        return out
    caption_hint = _visio_caption_from_paragraph(para)
    seen_rel: set = set()

    for rel_id in _image_rel_ids_from_paragraph(para):
        if rel_id in seen_rel:
            continue
        seen_rel.add(rel_id)
        blob, ctype = _part_blob_from_rel_id(doc, rel_id)
        if not blob:
            continue
        norm_blob, norm_ctype = _resolve_figure_blob(blob, ctype, work_dir=work_dir)
        if not norm_blob or not norm_ctype:
            continue
        out.append((norm_blob, norm_ctype, caption_hint))

    return out


def _append_figures_from_paragraph(
    para: Paragraph,
    doc: DocxDocument,
    *,
    work_dir: str,
    images_dir: str = "",
    figure_idx: int,
    html_parts: List[str],
    md_parts: List[str],
    figures: List[Dict[str, Any]],
) -> int:
    caption_hint = _visio_caption_from_paragraph(para)
    rel_ids = _image_rel_ids_from_paragraph(para)
    seen_rel: set = set()
    handled = False

    for rel_id in rel_ids:
        if rel_id in seen_rel:
            continue
        seen_rel.add(rel_id)
        blob, ctype = _part_blob_from_rel_id(doc, rel_id)
        if not blob:
            continue
        handled = True
        norm_blob, norm_ctype = _resolve_figure_blob(blob, ctype, work_dir=work_dir)
        figure_idx += 1
        cap = caption_hint or f"图{figure_idx}"
        if norm_blob and norm_ctype:
            fig_html, meta = _figure_html(norm_blob, norm_ctype, figure_idx, caption=cap, images_dir=images_dir)
            html_parts.append(fig_html)
            md_parts.append(_figure_markdown(norm_blob, norm_ctype, figure_idx, caption=cap))
            if caption_hint:
                meta["source"] = "visio_ole"
            figures.append(meta)
        else:
            reason = "矢量图预览未能转为 PNG，请确认已安装 LibreOffice 与 ImageMagick 后重新上传"
            fig_html, meta = _figure_placeholder_html(figure_idx, caption=cap, reason=reason)
            html_parts.append(fig_html)
            md_parts.append(f"【{cap}】（预览未转换）")
            meta["source"] = "visio_ole"
            meta["rel_id"] = rel_id
            figures.append(meta)
            logger.info(
                "[word_native] figure %d rel=%s ctype=%s len=%d convert failed",
                figure_idx, rel_id, ctype, len(blob),
            )

    if not handled and caption_hint:
        figure_idx += 1
        cap = caption_hint or "Visio图"
        fig_html, meta = _figure_placeholder_html(
            figure_idx,
            caption=cap,
            reason="矢量图预览未能转为 PNG，请确认已安装 LibreOffice 或 ImageMagick 后重新上传",
        )
        html_parts.append(fig_html)
        md_parts.append(f"【{cap}】（预览未转换）")
        meta["source"] = "visio_ole"
        figures.append(meta)
    return figure_idx


def _figure_html(blob: bytes, content_type: str, index: int, caption: str = "", images_dir: str = "") -> Tuple[str, Dict[str, Any]]:
    mime = _normalize_browser_image_mime(content_type, blob)
    if not mime or not blob:
        raise ValueError("non-displayable figure blob")
    b64 = base64.b64encode(blob).decode("ascii")
    cap = html.escape(caption) if caption else ""
    fig = (
        f'<figure class="word-native-figure" data-figure-index="{index}">'
        f'<img src="data:{mime};base64,{b64}" alt="图{index}" style="max-width:100%;height:auto;"/>'
        f"{f'<figcaption>{cap}</figcaption>' if cap else ''}"
        f"</figure>"
    )
    meta = {
        "figure_index": index,
        "page": 1,
        "content_type": mime,
        "blob": blob,
        "caption": caption,
    }
    if images_dir:
        try:
            os.makedirs(images_dir, exist_ok=True)
            ext = ".png" if "png" in mime else ".jpg"
            fpath = os.path.join(images_dir, f"fig_{index:04d}{ext}")
            with open(fpath, "wb") as f:
                f.write(blob)
            meta["path"] = fpath
            meta["file"] = fpath
        except OSError as exc:
            logger.debug("[word_native] figure file write failed idx=%s: %s", index, exc)
    return fig, meta


def _figure_placeholder_html(index: int, caption: str = "", reason: str = "") -> Tuple[str, Dict[str, Any]]:
    cap = html.escape(caption or f"图{index}")
    hint = html.escape(reason or "Visio/矢量图预览未能转换，请安装 LibreOffice 或 ImageMagick 后重新上传")
    fig = (
        f'<figure class="word-native-figure word-figure-placeholder" data-figure-index="{index}">'
        f'<p style="margin:8px 0;padding:8px 12px;background:#fff7e6;border-left:3px solid #faad14;">'
        f"<em>【{cap}】</em> {hint}</p></figure>"
    )
    return fig, {
        "figure_index": index,
        "page": 1,
        "content_type": "text/html",
        "caption": caption or f"图{index}",
        "placeholder": True,
    }


def parse_word_document(
    file_path: str,
    *,
    ocr_output_dir: Optional[str] = None,
) -> Dict[str, Any]:
    """
    解析 Word，返回与 GetTestMethod parser 兼容的结构，并附带 full_html。
    """
    if docx is None:
        raise RuntimeError("缺少 python-docx，请安装：pip install python-docx")

    work_dir = ocr_output_dir or tempfile.mkdtemp(prefix="word_native_")
    images_dir = os.path.join(work_dir, "images")
    os.makedirs(images_dir, exist_ok=True)
    docx_path, converted = _ensure_docx_path(file_path, work_dir)
    doc = docx.Document(docx_path)
    numbering = _WordNumberingResolver(doc)

    text_parts: List[str] = []
    html_parts: List[str] = []
    md_parts: List[str] = []
    tables_html: List[str] = []
    all_tables: List[Dict[str, Any]] = []
    figures: List[Dict[str, Any]] = []
    figure_idx = 0
    paragraph_map: List[Dict[str, Any]] = []
    char_pos = 0
    para_index = 0

    for block in _iter_body_blocks(doc):
        if isinstance(block, Paragraph):
            display = _display_text_for_paragraph(block, numbering)
            if not display:
                figure_idx = _append_figures_from_paragraph(
                    block, doc,
                    work_dir=work_dir,
                    images_dir=images_dir,
                    figure_idx=figure_idx,
                    html_parts=html_parts,
                    md_parts=md_parts,
                    figures=figures,
                )
                if _image_parts_from_paragraph(block, doc, work_dir=work_dir):
                    para_index += 1
                continue

            m = NUMBERED_HEADING_RE.match(display)
            section_no = m.group(1) if m else ""
            level = detect_paragraph_heading_level(block, display)
            if not level and section_no:
                level = min(max(len(section_no.split(".")), 1), 6)
            has_num, num_id, ilvl = _paragraph_numpr_info(block)
            char_start = char_pos
            char_end = char_start + len(display)
            paragraph_map.append({
                "para_index": para_index,
                "char_start": char_start,
                "char_end": char_end,
                "text": display,
                "raw_text": (block.text or "").strip(),
                "is_heading": bool(level),
                "level": int(level or 0),
                "section_no": section_no,
                "has_numpr": has_num,
                "num_id": num_id,
                "ilvl": ilvl,
            })
            para_index += 1

            if level:
                aid = _section_anchor_id(section_no) if section_no else f"wjgl-sec-pos-{char_start}"
                anchor = f' id="{html.escape(aid, quote=True)}"'
                data_no = f' data-section-no="{html.escape(section_no)}"' if section_no else ""
                html_parts.append(
                    f'<h{level}{anchor} class="wjgl-section-heading"{data_no}>'
                    f"<strong>{html.escape(display)}</strong></h{level}>"
                )
                hmarks = "#" * min(max(int(level) + 1, 2), 6)
                md_parts.append(f"{hmarks} {display}")
            else:
                inner = _runs_to_html(block)
                prefix = numbering.prefix_for_paragraph(block)
                if prefix and not (block.text or "").strip().startswith(prefix):
                    inner = f"{html.escape(prefix)} {inner}".strip()
                if inner:
                    html_parts.append(f"<p>{inner}</p>")
                md_body = _runs_to_markdown(block)
                if prefix and not (block.text or "").strip().startswith(prefix):
                    md_body = f"{prefix} {md_body}".strip() if md_body else display
                elif not md_body:
                    md_body = display
                if md_body:
                    md_parts.append(md_body)

            text_parts.append(display)
            char_pos = char_end + 2

            figure_idx = _append_figures_from_paragraph(
                block, doc,
                work_dir=work_dir,
                images_dir=images_dir,
                figure_idx=figure_idx,
                html_parts=html_parts,
                md_parts=md_parts,
                figures=figures,
            )

        elif isinstance(block, Table):
            tbl_html = _table_to_html(block)
            tbl_gfm = _table_to_gfm(block)
            if tbl_html:
                tables_html.append(tbl_html)
                all_tables.append({"page": 1, "html": tbl_html})
                html_parts.append(tbl_html)
            if tbl_gfm:
                md_parts.append(tbl_gfm)
            rows_text = []
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    rows_text.append(" | ".join(cells))
            if rows_text:
                tbl_text = "【表格】\n" + "\n".join(rows_text)
                char_start = char_pos
                char_end = char_start + len(tbl_text)
                paragraph_map.append({
                    "para_index": para_index,
                    "char_start": char_start,
                    "char_end": char_end,
                    "text": tbl_text,
                    "raw_text": tbl_text,
                    "is_heading": False,
                    "level": 0,
                    "section_no": "",
                    "is_table": True,
                })
                para_index += 1
                text_parts.append(tbl_text)
                char_pos = char_end + 2
            for cell_para in _paragraphs_in_table(block):
                figure_idx = _append_figures_from_paragraph(
                    cell_para, doc,
                    work_dir=work_dir,
                    images_dir=images_dir,
                    figure_idx=figure_idx,
                    html_parts=html_parts,
                    md_parts=md_parts,
                    figures=figures,
                )

    full_text = "\n\n".join(text_parts).strip()
    full_html = "".join(html_parts)
    full_markdown = "\n\n".join(p for p in md_parts if p and str(p).strip()).strip()

    if not full_text and full_html:
        full_text = re.sub(r"<[^>]+>", " ", full_html)
        full_text = re.sub(r"\s+", " ", full_text).strip()

    logger.info(
        "[word_native] 解析完成 file=%s converted=%s text_len=%d md_len=%d tables=%d figures=%d html_len=%d paragraphs=%d",
        os.path.basename(file_path),
        converted,
        len(full_text),
        len(full_markdown),
        len(tables_html),
        len(figures),
        len(full_html),
        len(paragraph_map),
    )

    page = {
        "page_number": 1,
        "text": full_text,
        "tables_html": tables_html,
        "figures": [
            {k: v for k, v in fig.items() if k != "blob"}
            for fig in figures
        ],
    }

    return {
        "mode": "docx_native",
        "parse_mode": "docx_native",
        "full_text": full_text,
        "full_html": full_html,
        "full_markdown": full_markdown,
        "pages": [page],
        "all_tables": all_tables,
        "figures_raw": figures,
        "ocr_used": False,
        "converted_from_doc": converted,
        "source_docx_path": docx_path if converted else None,
        "word_paragraph_map": paragraph_map,
        "ocr_output_dir": work_dir,
        "images_dir": images_dir,
    }


def extract_word_headings_merged(docx_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """合并多种策略的 Word 标题识别，提高覆盖率。"""
    if docx is None:
        return [], "error"
    try:
        docx_path, _ = _ensure_docx_path(docx_path)
        doc = docx.Document(docx_path)
    except Exception as exc:
        logger.warning("[word_native] 打开 Word 失败: %s", exc)
        return [], "error"

    from .document_heading_extract import (
        build_toc_tree,
        extract_headings_by_number,
        extract_headings_by_style,
        parse_section_number,
    )

    merged: List[Dict[str, Any]] = []
    seen: set = set()

    def _add(title: str, section_no: str, level: int, source: str) -> None:
        title = (title or "").strip()
        if not title:
            return
        if not section_no:
            section_no, _ = parse_section_number(title)
        key = (section_no or "", title[:80])
        if key in seen:
            return
        seen.add(key)
        merged.append({
            "title": title,
            "sectionNo": section_no or "",
            "level": min(max(int(level or 1), 1), 6),
            "_source": source,
        })

    for item in extract_headings_by_style(doc):
        _add(item.get("title", ""), item.get("sectionNo", ""), item.get("level", 1), "style")

    for block in _iter_body_blocks(doc):
        if not isinstance(block, Paragraph):
            continue
        text = (block.text or "").strip()
        if not text:
            continue
        lv = detect_paragraph_heading_level(block)
        if lv:
            sn, _ = parse_section_number(text)
            _add(text, sn or "", lv, "body")

    for item in extract_headings_by_number(doc):
        _add(item.get("title", ""), item.get("sectionNo", ""), item.get("level", 1), "number")

    if not merged:
        return [], "empty"
    method = "word_merged"
    if len(merged) >= 3:
        logger.info("[word_native] 合并标题 %d 条", len(merged))
    return build_toc_tree(merged), method
