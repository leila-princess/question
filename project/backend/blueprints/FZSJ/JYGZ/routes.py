"""
JYGZ 基于规则的试验方案生成蓝图
"""
import json
import os
import re
import uuid
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional
from flask import Blueprint, request, jsonify, send_file, Response
from config import Config
from utils.neo4j_flgl_status import flgl_status_cypher_param
from blueprints.FZSJ.JYLLM.method_synr_markdown import (
    apply_method_heading_prefix,
    strip_redundant_method_name_heading,
    synr_html_to_markdown as jyllm_synr_html_to_markdown,
)
from blueprints.FZSJ.JYLLM.project_tree_utils import rebuild_indicator_forest, rebuild_project_forest

logger = logging.getLogger(__name__)

jygz_bp = Blueprint('jygz', __name__, url_prefix='/api/fzsj/jygz')

# 路径配置
BASE_DIR = Path(__file__).parent
TEMPLATES_DIR = BASE_DIR / 'templates'
CONFIGS_DIR = BASE_DIR / 'configs'
TEMPLATES_DIR.mkdir(exist_ok=True)
CONFIGS_DIR.mkdir(exist_ok=True)

# 页面标识符，用于区分基于规则和基于LLM的历史记录
PAGE_TYPE = 'jygz'  # 规则型
LLM_PAGE_TYPE = 'jyllm'  # LLM型

JYGZ_HISTORY_DDL = """
CREATE TABLE IF NOT EXISTS jygz_history (
    id BIGINT AUTO_INCREMENT PRIMARY KEY COMMENT '主键ID',
    task_id VARCHAR(100) NOT NULL COMMENT '任务唯一标识',
    user_id VARCHAR(64) NOT NULL COMMENT '用户ID',
    scheme_name VARCHAR(200) COMMENT '方案名称',
    scheme_description TEXT COMMENT '方案描述',
    selected_equipment JSON COMMENT '选择的设备列表',
    selected_environments JSON COMMENT '选择的环境列表',
    selected_projects JSON COMMENT '选择的项目列表',
    selected_indicators JSON COMMENT '选择的指标列表',
    selected_methods JSON COMMENT '选择的方法列表',
    selected_files JSON COMMENT '参考文件列表',
    environment_config LONGTEXT COMMENT '环境详细配置',
    template_id VARCHAR(64) COMMENT '使用的模板ID',
    template_name VARCHAR(200) COMMENT '模板名称',
    config_snapshot LONGTEXT COMMENT '完整前端状态 JSON',
    generated_content LONGTEXT COMMENT '最终方案 Markdown',
    word_count INT DEFAULT 0,
    generation_status VARCHAR(20) DEFAULT 'complete',
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
    completed_at TIMESTAMP NULL,
    UNIQUE KEY uk_jygz_task (task_id),
    INDEX idx_jygz_user (user_id),
    INDEX idx_jygz_updated (updated_at)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
COMMENT='基于规则 JYGZ 方案历史'
"""


def _ensure_jygz_history_table(conn):
    with conn.cursor() as c:
        c.execute(JYGZ_HISTORY_DDL)


def _derive_jygz_selections_from_config(config: dict) -> dict:
    out = {
        'selected_equipment': [],
        'selected_environments': [],
        'selected_projects': [],
        'selected_indicators': [],
        'selected_methods': [],
    }
    if not config or not isinstance(config, dict):
        return out
    for x in config.get('equipment_list') or []:
        if isinstance(x, dict):
            out['selected_equipment'].append({
                'id': str(x.get('id', '')),
                'name': str(x.get('name', '')),
                'type': str(x.get('type', '')),
            })
    for e in config.get('environments') or []:
        if isinstance(e, dict):
            out['selected_environments'].append({
                'env_id': str(e.get('env_id', '')),
                'env_name': str(e.get('env_name', '')),
            })
    pids = config.get('selected_projects') or []
    for pid in pids:
        pd = next(
            (p for p in (config.get('projects_data') or []) if str(p.get('id', '')) == str(pid)),
            None,
        )
        if isinstance(pd, dict):
            out['selected_projects'].append({
                'id': str(pd.get('id', '')),
                'name': str(pd.get('name', '')),
            })
    for iid in (config.get('selected_indicators') or []):
        ind = next(
            (p for p in (config.get('indicators_data') or []) if str(p.get('id', '')) == str(iid)),
            None,
        )
        if isinstance(ind, dict):
            out['selected_indicators'].append({
                'id': str(ind.get('id', '')),
                'name': str(ind.get('name', '')),
            })
    for mid in (config.get('selected_methods') or []):
        md = next(
            (m for m in (config.get('methods_data') or []) if str(m.get('id', '')) == str(mid)),
            None,
        )
        if isinstance(md, dict):
            out['selected_methods'].append({
                'id': str(md.get('id', '')),
                'name': str(md.get('name', '')),
            })
    return out


# ==================== Pydantic 模型（用 dict 验证代替） ====================

def validate_config(data: dict) -> tuple[bool, str]:
    """验证配置数据结构"""
    required_fields = ['设备名称', '编号', '用途', '环境描述', '任务来源']
    for field in required_fields:
        if field not in data:
            return False, f'缺少必填字段: {field}'
    
    if '环境列表' not in data or not isinstance(data['环境列表'], list):
        return False, '环境列表格式错误'
    
    if '边界性能指标' not in data or not isinstance(data['边界性能指标'], list):
        return False, '边界性能指标格式错误'
    
    return True, ''


# ==================== 模板管理 ====================

@jygz_bp.route('/templates/list', methods=['GET'])
def list_templates_list():
    user_id = request.args.get('user_id', 'admin')
    try:
        templates = []
        for f in TEMPLATES_DIR.glob('*.md'):
            templates.append({
                'id': f.stem,
                'name': f.stem,
                'description': f'内置模板: {f.stem}',
                'category': '内置',
                'version': '1.0'
            })
        
        # 同时从Chroma获取用户保存的模板
        if Config.CHROMA_CLIENT:
            try:
                collections = Config.CHROMA_CLIENT.list_collections()
                template_collection = None
                for col in collections:
                    if col.name == 'jygz_templates':
                        template_collection = col
                        break
                
                if template_collection:
                    results = template_collection.get(limit=100, include=['metadatas', 'documents'])
                    for i, meta in enumerate(results.get('metadatas', [])):
                        if meta:
                            templates.append({
                                'id': meta.get('template_name', f'user_template_{i}'),
                                'name': meta.get('template_name', f'user_template_{i}'),
                                'description': meta.get('description', ''),
                                'category': '用户模板',
                                'version': '1.0'
                            })
            except Exception as e:
                logger.warning(f'从Chroma获取模板失败: {e}')
        
        return jsonify({'success': True, 'templates': templates})
    except Exception as e:
        logger.error(f'获取模板列表失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/templates', methods=['GET'])
def list_templates():
    """返回可用的模板列表"""
    try:
        templates = []
        for f in TEMPLATES_DIR.glob('*.md'):
            templates.append({
                'name': f.stem,
                'path': str(f),
                'modified': datetime.fromtimestamp(f.stat().st_mtime).isoformat(),
                'size': f.stat().st_size
            })
        
        # 同时从Chroma获取用户保存的模板
        if Config.CHROMA_CLIENT:
            try:
                collections = Config.CHROMA_CLIENT.list_collections()
                template_collection = None
                for col in collections:
                    if col.name == 'jygz_templates':
                        template_collection = col
                        break
                
                if template_collection:
                    # 只获取当前页面类型的模板
                    results = template_collection.get(
                        where={'page_type': PAGE_TYPE}, 
                        limit=100, 
                        include=['metadatas', 'documents']
                    )
                    for i, meta in enumerate(results.get('metadatas', [])):
                        if meta:
                            templates.append({
                                'name': meta.get('template_name', f'user_template_{i}'),
                                'path': 'chroma',
                                'modified': meta.get('created_at', ''),
                                'is_user_template': True,
                                'description': meta.get('description', '')
                            })
            except Exception as e:
                logger.warning(f'从Chroma获取模板失败: {e}')
        
        return jsonify({'success': True, 'data': templates})
    except Exception as e:
        logger.error(f'获取模板列表失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/templates/detail/<template_id>', methods=['GET'])
def get_template_detail(template_id: str):
    try:
        template_path = TEMPLATES_DIR / f'{template_id}.md'
        if template_path.exists():
            content = template_path.read_text(encoding='utf-8')
            return jsonify({
                'success': True,
                'template': {
                    'id': template_id,
                    'metadata': {
                        'name': template_id,
                        'description': f'内置模板: {template_id}',
                        'category': '内置'
                    },
                    'content': content
                }
            })
        
        # 从Chroma查找
        if Config.CHROMA_CLIENT:
            try:
                collections = Config.CHROMA_CLIENT.list_collections()
                for col in collections:
                    if col.name == 'jygz_templates':
                        results = col.get(where={'template_name': template_id}, limit=1, include=['documents', 'metadatas'])
                        if results.get('documents'):
                            meta = results['metadatas'][0] if results.get('metadatas') else {}
                            return jsonify({
                                'success': True,
                                'template': {
                                    'id': template_id,
                                    'metadata': {
                                        'name': template_id,
                                        'description': meta.get('description', ''),
                                        'category': '用户模板'
                                    },
                                    'content': results['documents'][0]
                                }
                            })
            except Exception as e:
                logger.warning(f'从Chroma获取模板失败: {e}')
        
        return jsonify({'success': False, 'error': '模板不存在'}), 404
    except Exception as e:
        logger.error(f'获取模板失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/templates/<name>', methods=['GET'])
def get_template(name: str):
    """获取指定模板内容"""
    try:
        template_path = TEMPLATES_DIR / f'{name}.md'
        if template_path.exists():
            content = template_path.read_text(encoding='utf-8')
            return jsonify({'success': True, 'data': {'name': name, 'content': content}})
        
        # 从Chroma查找
        if Config.CHROMA_CLIENT:
            try:
                collections = Config.CHROMA_CLIENT.list_collections()
                for col in collections:
                    if col.name == 'jygz_templates':
                        results = col.get(where={'template_name': name}, limit=1, include=['documents', 'metadatas'])
                        if results.get('documents'):
                            return jsonify({
                                'success': True,
                                'data': {
                                    'name': name,
                                    'content': results['documents'][0],
                                    'is_user_template': True
                                }
                            })
            except Exception as e:
                logger.warning(f'从Chroma获取模板失败: {e}')
        
        return jsonify({'success': False, 'error': '模板不存在'}), 404
    except Exception as e:
        logger.error(f'获取模板失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/templates/save', methods=['POST'])
def save_template_api():
    data = request.get_json()
    template_name = data.get('name')
    content = data.get('content')
    description = data.get('description', '')
    page_type = data.get('page_type', PAGE_TYPE)  # 默认为当前页面类型
    
    if not template_name or not content:
        return jsonify({'success': False, 'error': '模板名称和内容不能为空'}), 400
    
    try:
        if Config.CHROMA_CLIENT:
            collection = Config.CHROMA_CLIENT.get_or_create_collection(name='jygz_templates')
            
            existing = collection.get(where={'template_name': template_name}, limit=1)
            if existing.get('ids'):
                collection.update(
                    ids=existing['ids'][0],
                    documents=[content],
                    metadatas=[{
                        'template_name': template_name,
                        'description': description,
                        'updated_at': datetime.now().isoformat(),
                        'page_type': page_type
                    }]
                )
            else:
                collection.add(
                    ids=[f'template_{template_name}_{datetime.now().timestamp()}'],
                    documents=[content],
                    metadatas=[{
                        'template_name': template_name,
                        'description': description,
                        'created_at': datetime.now().isoformat(),
                        'page_type': page_type
                    }]
                )
            
            return jsonify({'success': True, 'message': '模板保存成功'})
        else:
            return jsonify({'success': False, 'error': 'Chroma未连接'}), 500
    except Exception as e:
        logger.error(f'保存模板失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/templates', methods=['POST'])
def save_template():
    """保存模板到Chroma"""
    data = request.get_json()
    template_name = data.get('name')
    content = data.get('content')
    description = data.get('description', '')
    
    if not template_name or not content:
        return jsonify({'success': False, 'error': '模板名称和内容不能为空'}), 400
    
    try:
        if Config.CHROMA_CLIENT:
            collection = Config.CHROMA_CLIENT.get_or_create_collection(name='jygz_templates')
            
            # 检查是否已存在
            existing = collection.get(where={'template_name': template_name}, limit=1)
            if existing.get('ids'):
                collection.update(
                    ids=existing['ids'][0],
                    documents=[content],
                    metadatas=[{
                        'template_name': template_name,
                        'description': description,
                        'updated_at': datetime.now().isoformat(),
                        'page_type': PAGE_TYPE
                    }]
                )
            else:
                collection.add(
                    ids=[f'template_{template_name}_{datetime.now().timestamp()}'],
                    documents=[content],
                    metadatas=[{
                        'template_name': template_name,
                        'description': description,
                        'created_at': datetime.now().isoformat(),
                        'page_type': PAGE_TYPE
                    }]
                )
            
            return jsonify({'success': True, 'message': '模板保存成功'})
        else:
            return jsonify({'success': False, 'error': 'Chroma未连接'}), 500
    except Exception as e:
        logger.error(f'保存模板失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/templates/delete/<template_name>', methods=['DELETE'])
def delete_template(template_name: str):
    """删除模板"""
    user_id = request.args.get('user_id', 'admin')
    
    try:
        if Config.CHROMA_CLIENT:
            collection = Config.CHROMA_CLIENT.get_or_create_collection(name='jygz_templates')
            
            # 检查模板是否存在
            existing = collection.get(where={'template_name': template_name}, limit=1)
            if existing.get('ids'):
                collection.delete(ids=existing['ids'][0])
                return jsonify({'success': True, 'message': '模板删除成功'})
            else:
                return jsonify({'success': False, 'error': '模板不存在'}), 404
        else:
            return jsonify({'success': False, 'error': 'Chroma未连接'}), 500
    except Exception as e:
        logger.error(f'删除模板失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


# ==================== 渲染 API ====================

def split_basis_standard_items(text: str) -> list[str]:
    """编制依据：按换行、顿号、中英文分号拆成多条（表格一行一条）。"""
    if not text or not str(text).strip():
        return []
    s = str(text).strip().replace('\r\n', '\n').replace('\r', '\n')
    parts = re.split(r'[\n、；;]+', s)
    return [p.strip() for p in parts if p.strip()]


def _md_table_cell(text) -> str:
    """避免 GFM 表格单元格内的 |、换行破坏表格解析。"""
    if text is None:
        return ''
    t = str(text).replace('\r\n', '\n').replace('\r', '\n')
    t = t.replace('|', '｜').replace('\n', ' ')
    return t.strip()


def _compact_gfm_table_row_gaps(md: str) -> str:
    """去掉以 | 开头的表格行之间的多余空行（编制依据等表格行紧密相连）。"""
    if not md:
        return md
    return re.sub(r'(\|[^\n]+\|)\n\n+(?=\|)', r'\1\n', md)


def _strip_legacy_scheme_layout_divs(md: str) -> str:
    """去掉历史模板里仅作布局的 div，避免 Markdown 预览/导出里出现 HTML 壳标签。"""
    if not md:
        return md
    md = re.sub(
        r'<div\s+class=["\']env-settings["\']\s*>\s*',
        '',
        md,
        flags=re.IGNORECASE,
    )
    # 闭合标签后接下一小节标题（###～######），含无项目时直接接「### 2.2」等情况
    md = re.sub(r'(?m)^</div>\s*\n(?=\s*#{3,6}\s)', '', md)
    return md


@jygz_bp.route('/render', methods=['POST'])
def render_template():
    try:
        from jinja2 import Template
        
        data = request.get_json()
        if not data:
            return jsonify({'success': False, 'error': '请求体为空'}), 400
        
        equipment_name = data.get('equipment_name', '')
        doc_number = data.get('doc_number', '')
        task_source = data.get('task_source', '')
        basis_standard = data.get('basis_standard', '')
        
        # 1. 处理依据标准（编制依据：换行、顿号、中英文分号 → 表格多行）
        standard_list = []
        if basis_standard:
            for p in split_basis_standard_items(basis_standard):
                standard_list.append({'name': _md_table_cell(p), 'source': '编制依据'})

        # 前端可将多环境合并为单节「试验环境」传入；模板见 simplified_template.md（环境参数字段默认不输出正文）
        env_list = []
        environments_data = data.get('environments', [])
        all_methods = data.get('methods', [])

        def _method_fields(method_data: dict, parent_prefix: str = "") -> dict:
            synr_raw = (
                method_data.get('SYNR')
                or method_data.get('synr')
                or ''
            )
            mid = str(method_data.get('id') or method_data.get('methodId') or '').strip() or 'unknown_method'
            imageurls = (
                method_data.get('imageurls')
                or method_data.get('image_urls')
                or method_data.get('ImageUrls')
            )
            files_raw = method_data.get('files')
            synr = jyllm_synr_html_to_markdown(synr_raw, mid, imageurls, files_raw)
            synr = strip_redundant_method_name_heading(synr, method_data.get('name', ''))
            synr = apply_method_heading_prefix(synr, parent_prefix)
            return {
                'name': method_data.get('name', '未命名方法'),
                'SYNR': synr,
            }

        def _nid(v):
            return str(v or '').strip()

        def _is_method_node(node: dict) -> bool:
            if not isinstance(node, dict):
                return False
            return any(node.get(k) for k in ('SYNR', 'synr', 'SYTJ', 'sytj', 'SYGC', 'sygc', 'JGPD', 'jgpd'))

        def _pick_methods_by_project(proj_id: str, proj_name: str) -> list:
            candidate_ids = {_nid(proj_id)}
            for m in all_methods:
                links = m.get('link_ids') or []
                for lid in links:
                    candidate_ids.add(_nid(lid))
            candidate_ids.discard('')
            methods_from_proj = []
            for m in all_methods:
                hit = False
                rel = m.get('related_projects') or m.get('relatedProjects') or []
                for r in rel:
                    rid = _nid(r.get('id'))
                    if rid and rid in candidate_ids:
                        hit = True
                        break
                if not hit:
                    linked = m.get('linkedProjectIds') or m.get('linked_project_ids') or []
                    if linked and candidate_ids.intersection({_nid(x) for x in linked}):
                        hit = True
                if not hit and proj_id:
                    if _nid(m.get('project_id')) == _nid(proj_id):
                        hit = True
                if not hit and proj_name:
                    m_name = str(m.get('name', ''))
                    m_label = str(m.get('label', ''))
                    m_desc = str(m.get('description', ''))
                    if proj_name in m_name or proj_name in m_label or proj_name in m_desc:
                        hit = True
                if hit:
                    methods_from_proj.append(m)
            return methods_from_proj

        def _collect_indicator_descendant_ids(node: dict) -> set:
            ids = set()
            nid = str((node or {}).get('id', '') or '').strip()
            if nid:
                ids.add(nid)
            for ch in (node or {}).get('children') or []:
                if isinstance(ch, dict):
                    ids.update(_collect_indicator_descendant_ids(ch))
            return ids

        def _pick_methods_by_indicator(ind_id: str, ind_name: str, descendant_ids: Optional[set] = None) -> list:
            target_ids = {str(x).strip() for x in (descendant_ids or set()) if str(x).strip()}
            if ind_id:
                target_ids.add(ind_id)
            picked = []
            for m in all_methods:
                hit = False
                rel = m.get('related_indicators') or m.get('relatedIndicators') or []
                for r in rel:
                    rid = str(r.get('id', '') or '').strip()
                    if rid and (rid in target_ids):
                        hit = True
                        break
                if not hit:
                    linked = m.get('linkedIndicatorIds') or m.get('linked_indicator_ids') or []
                    if target_ids.intersection({str(x).strip() for x in linked}):
                        hit = True
                if not hit and ind_name:
                    m_name = str(m.get('name', ''))
                    m_label = str(m.get('label', ''))
                    m_desc = str(m.get('description', ''))
                    if ind_name in m_name or ind_name in m_label or ind_name in m_desc:
                        hit = True
                if hit:
                    picked.append(m)
            return picked

        def _walk_project_nodes(nodes: list, parent_no: str) -> list:
            out: list = []
            proj_seq = 0
            for node in (nodes or []):
                if not isinstance(node, dict):
                    continue
                if _is_method_node(node):
                    continue
                proj_seq += 1
                proj_no = f"{parent_no}.{proj_seq}"
                proj_name = node.get('name', '未知项目')
                proj_desc = node.get('description', '')
                proj_id = str(node.get('id', '') or '')

                embedded_methods = node.get('methods') or []
                child_nodes = node.get('children') or []
                child_projects = [
                    ch for ch in child_nodes
                    if isinstance(ch, dict) and not _is_method_node(ch)
                ]
                if not embedded_methods and child_nodes and not child_projects:
                    if all(_is_method_node(ch) for ch in child_nodes if isinstance(ch, dict)):
                        embedded_methods = child_nodes
                methods_data = embedded_methods or _pick_methods_by_project(proj_id, proj_name)

                methods_for_template = []
                for method_data in methods_data:
                    f = _method_fields(method_data, proj_no)
                    methods_for_template.append(f)

                child_blocks = _walk_project_nodes(child_projects, proj_no) if child_projects else []
                out.append({
                    'name': proj_name,
                    'description': proj_desc,
                    'methods': methods_for_template,
                    'number': proj_no,
                    'children': child_blocks,
                })
            return out

        def _walk_indicator_nodes(nodes: list, parent_no: str) -> list:
            out: list = []
            ind_seq = 0
            for node in (nodes or []):
                if not isinstance(node, dict):
                    continue
                if _is_method_node(node):
                    continue
                ind_seq += 1
                ind_no = f"{parent_no}.{ind_seq}"
                ind_name = node.get('name', '未知指标')
                ind_desc = node.get('description', '')
                ind_id_str = str(node.get('id', '') or '').strip()

                embedded_methods = node.get('methods') or []
                child_nodes = node.get('children') or []
                child_inds = [
                    ch for ch in child_nodes
                    if isinstance(ch, dict) and not _is_method_node(ch)
                ]
                if not embedded_methods and child_nodes and not child_inds:
                    if all(_is_method_node(ch) for ch in child_nodes if isinstance(ch, dict)):
                        embedded_methods = child_nodes
                descendant_ids = _collect_indicator_descendant_ids(node)
                methods_data = embedded_methods or _pick_methods_by_indicator(ind_id_str, ind_name, descendant_ids)

                methods_for_template = []
                for method_data in methods_data:
                    f = _method_fields(method_data, ind_no)
                    methods_for_template.append(f)

                child_blocks = _walk_indicator_nodes(child_inds, ind_no) if child_inds else []
                out.append({
                    'name': ind_name,
                    'description': ind_desc,
                    'methods': methods_for_template,
                    'number': ind_no,
                    'children': child_blocks,
                })
            return out

        for env_idx, env_data in enumerate(environments_data, 1):
            env_name = env_data.get('env_name', '未知环境')
            settings = env_data.get('settings', {})

            projects_data = rebuild_project_forest(env_data.get('projects', []))
            projects_for_template = _walk_project_nodes(projects_data, f"2.1.{env_idx}")

            env_list.append({
                'name': env_name,
                'number': f"2.1.{env_idx}",
                'settings': settings,
                'projects': projects_for_template
            })

        indicators_list_for_template = []
        indicators_data = rebuild_indicator_forest(data.get('indicators', []))
        
        if not indicators_data:
            indicators_list_for_template.append({
                'name': '边界性能指标（未配置）',
                'description': '请在基础配置中选择指标以生成此部分内容',
                'methods': [],
                'children': [],
                'number': '2.2.1'
            })
        else:
            indicators_list_for_template = _walk_indicator_nodes(indicators_data, "2.2")

        template_data = {
            '设备名称': equipment_name,
            '编号': doc_number,
            '任务来源': task_source,
            '术语列表': [],
            '依据标准列表': standard_list,
            '环境列表': env_list,
            '边界性能指标': indicators_list_for_template
        }
        
        template_name = data.get('_template_name', 'simplified_template')
        template_path = TEMPLATES_DIR / f'{template_name}.md'
        
        if template_path.exists():
            template_content = template_path.read_text(encoding='utf-8')
        else:
            template_content = get_builtin_default_template()
        
        template = Template(template_content)
        rendered = template.render(**template_data)
        rendered = _strip_legacy_scheme_layout_divs(_compact_gfm_table_row_gaps(rendered))

        return jsonify({'success': True, 'data': rendered})
    except Exception as e:
        logger.error(f'渲染模板失败: {e}')
        import traceback
        logger.error(traceback.format_exc())
        return jsonify({'success': False, 'error': str(e)}), 500






# ==================== 导出 API ====================

@jygz_bp.route('/export/pdf', methods=['POST'])
def export_pdf():
    """
    优先：请求体带 markdown（与前端预览一致），使用与 JYLLM export PDF 相同的 scheme_pdf_export。
    否则：旧版中文配置字段 + default_template 渲染后再转 PDF。
    """
    try:
        data = request.get_json()
        if not data:
            return jsonify({'success': False, 'error': '请求体为空'}), 400

        md = (data.get('markdown') or data.get('rendered_markdown') or '').strip()
        if md:
            from utils.scheme_pdf_export import export_markdown_to_pdf_bytes

            title = (data.get('title') or data.get('equipment_name') or '试验方案').strip() or '试验方案'
            doc_id = str(data.get('doc_number') or data.get('编号') or 'export').strip()
            pdf_content, pdf_err = export_markdown_to_pdf_bytes(md, doc_title=title, doc_id=doc_id)
            if pdf_content:
                fname = f"scheme_{doc_id or 'export'}.pdf"
                return Response(
                    pdf_content,
                    mimetype='application/pdf',
                    headers={'Content-Disposition': f'attachment; filename={fname}'},
                )
            logger.error('JYGZ 导出 PDF: %s', pdf_err)
            return jsonify({'success': False, 'error': pdf_err or 'PDF生成失败'}), 500

        from jinja2 import Template

        valid, error = validate_config(data)
        if not valid:
            return jsonify({'success': False, 'error': error}), 400

        template_name = data.get('_template_name', 'default_template')
        template_path = TEMPLATES_DIR / f'{template_name}.md'
        template_content = (
            template_path.read_text(encoding='utf-8') if template_path.exists() else get_default_template()
        )
        rendered = Template(template_content).render(**data)
        pdf_content = markdown_to_pdf(rendered, data.get('设备名称', '试验方案'))
        return Response(
            pdf_content,
            mimetype='application/pdf',
            headers={'Content-Disposition': f'attachment; filename=scheme_{data.get("编号", "export")}.pdf'},
        )
    except Exception as e:
        logger.error(f'导出PDF失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/export/word', methods=['POST'])
def export_word():
    """
    优先：请求体带 markdown，使用与 JYLLM export Word 相同的 utils.scheme_markdown_to_docx。
    否则：旧版中文配置 + default_template。
    """
    try:
        data = request.get_json()
        if not data:
            return jsonify({'success': False, 'error': '请求体为空'}), 400

        md = (data.get('markdown') or data.get('rendered_markdown') or '').strip()
        if md:
            try:
                from utils.scheme_markdown_to_docx import markdown_to_docx_bytes

                blob = markdown_to_docx_bytes(md)
                doc_id = str(data.get('doc_number') or data.get('编号') or 'export').strip()
                fname = f"scheme_{doc_id or 'export'}.docx"
                return Response(
                    blob,
                    mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                    headers={'Content-Disposition': f'attachment; filename={fname}'},
                )
            except ImportError:
                return jsonify({'success': False, 'error': 'Word导出需要安装python-docx'}), 500

        from jinja2 import Template
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.oxml.ns import qn
        import io
        import re

        valid, error = validate_config(data)
        if not valid:
            return jsonify({'success': False, 'error': error}), 400

        template_name = data.get('_template_name', 'default_template')
        template_path = TEMPLATES_DIR / f'{template_name}.md'
        template_content = (
            template_path.read_text(encoding='utf-8') if template_path.exists() else get_default_template()
        )
        rendered = Template(template_content).render(**data)

        doc = Document()
        style = doc.styles['Normal']
        font = style.font
        font.name = '微软雅黑'
        font.size = Pt(11)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

        for line in rendered.split('\n'):
            line_stripped = line.strip()
            if not line_stripped:
                doc.add_paragraph('')
            elif line_stripped.startswith('# '):
                doc.add_heading(line_stripped[2:], level=1)
            elif line_stripped.startswith('## '):
                doc.add_heading(line_stripped[3:], level=2)
            elif line_stripped.startswith('### '):
                doc.add_heading(line_stripped[4:], level=3)
            elif line_stripped.startswith('#### '):
                doc.add_heading(line_stripped[5:], level=4)
            elif line_stripped.startswith('##### '):
                doc.add_heading(line_stripped[6:], level=5)
            elif line_stripped.startswith('- ') or line_stripped.startswith('* '):
                doc.add_paragraph(line_stripped[2:], style='List Bullet')
            elif re.match(r'^\d+\.\s', line_stripped):
                doc.add_paragraph(line_stripped, style='List Number')
            else:
                doc.add_paragraph(line_stripped)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return Response(
            buffer.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            headers={'Content-Disposition': f'attachment; filename=scheme_{data.get("编号", "export")}.docx'},
        )
    except Exception as e:
        logger.error(f'导出Word失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


# ==================== 配置管理 ====================

@jygz_bp.route('/config/save', methods=['POST'])
def save_config():
    """保存用户配置到MySQL"""
    data = request.get_json()
    config_name = data.get('name')
    config_data = data.get('data')
    user_id = data.get('user_id', 'admin')
    
    if not config_name or not config_data:
        return jsonify({'success': False, 'error': '配置名称和数据不能为空'}), 400
    
    if not Config.MYSQL_POOL:
        return jsonify({'success': False, 'error': 'MySQL未连接'}), 500
    
    try:
        with Config.MYSQL_POOL.connection() as conn:
            with conn.cursor() as cursor:
                # 创建配置表（如果不存在）
                cursor.execute("""
                    CREATE TABLE IF NOT EXISTS jygz_configs (
                        id INT AUTO_INCREMENT PRIMARY KEY,
                        config_name VARCHAR(255) NOT NULL,
                        user_id VARCHAR(50) NOT NULL,
                        page_type VARCHAR(20) NOT NULL DEFAULT 'jygz',
                        config_data JSON NOT NULL,
                        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                        updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                        INDEX idx_user_page (user_id, page_type)
                    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
                """)
                
                # 检查是否已存在
                cursor.execute(
                    "SELECT id FROM jygz_configs WHERE config_name = %s AND user_id = %s AND page_type = %s",
                    (config_name, user_id, PAGE_TYPE)
                )
                existing = cursor.fetchone()
                
                if existing:
                    cursor.execute(
                        "UPDATE jygz_configs SET config_data = %s, updated_at = NOW() WHERE id = %s",
                        (json.dumps(config_data, ensure_ascii=False), existing[0])
                    )
                else:
                    cursor.execute(
                        "INSERT INTO jygz_configs (config_name, user_id, page_type, config_data) VALUES (%s, %s, %s, %s)",
                        (config_name, user_id, PAGE_TYPE, json.dumps(config_data, ensure_ascii=False))
                    )
                conn.commit()
        
        return jsonify({'success': True, 'message': '配置保存成功'})
    except Exception as e:
        logger.error(f'保存配置失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/config/load/<name>', methods=['GET'])
def load_config(name: str):
    """加载已保存的配置"""
    user_id = request.args.get('user_id', 'admin')
    
    if not Config.MYSQL_POOL:
        return jsonify({'success': False, 'error': 'MySQL未连接'}), 500
    
    try:
        with Config.MYSQL_POOL.connection() as conn:
            with conn.cursor() as cursor:
                cursor.execute(
                    "SELECT config_data FROM jygz_configs WHERE config_name = %s AND user_id = %s AND page_type = %s",
                    (name, user_id, PAGE_TYPE)
                )
                row = cursor.fetchone()
                
                if row:
                    return jsonify({'success': True, 'data': row[0] if isinstance(row[0], dict) else json.loads(row[0])})
                else:
                    return jsonify({'success': False, 'error': '配置不存在'}), 404
    except Exception as e:
        logger.error(f'加载配置失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/config/list', methods=['GET'])
def list_configs():
    """列出用户保存的配置"""
    user_id = request.args.get('user_id', 'admin')
    
    if not Config.MYSQL_POOL:
        return jsonify({'success': False, 'error': 'MySQL未连接'}), 500
    
    try:
        with Config.MYSQL_POOL.connection() as conn:
            with conn.cursor() as cursor:
                cursor.execute(
                    "SELECT config_name, created_at, updated_at FROM jygz_configs WHERE user_id = %s AND page_type = %s ORDER BY updated_at DESC",
                    (user_id, PAGE_TYPE)
                )
                rows = cursor.fetchall()
                configs = [
                    {'name': r[0], 'created_at': str(r[1]), 'updated_at': str(r[2])}
                    for r in rows
                ]
        
        return jsonify({'success': True, 'data': configs})
    except Exception as e:
        logger.error(f'列出配置失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/config/delete/<name>', methods=['DELETE'])
def delete_config(name: str):
    """删除配置"""
    user_id = request.args.get('user_id', 'admin')
    
    if not Config.MYSQL_POOL:
        return jsonify({'success': False, 'error': 'MySQL未连接'}), 500
    
    try:
        with Config.MYSQL_POOL.connection() as conn:
            with conn.cursor() as cursor:
                cursor.execute(
                    "DELETE FROM jygz_configs WHERE config_name = %s AND user_id = %s AND page_type = %s",
                    (name, user_id, PAGE_TYPE)
                )
                conn.commit()
        
        return jsonify({'success': True, 'message': '配置删除成功'})
    except Exception as e:
        logger.error(f'删除配置失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/history/save', methods=['POST'])
def save_jygz_history():
    """
    保存基于规则的完整方案到 jygz_history：方案名、全量 state（config + 环境编辑态、env_settings_data）、生成正文。
    """
    data = request.get_json() or {}
    user_id = (data.get('user_id') or 'admin').strip()
    scheme_name = (data.get('scheme_name') or '').strip()
    if not scheme_name:
        return jsonify({'success': False, 'error': '请填写方案名称'}), 400
    if not Config.MYSQL_POOL:
        return jsonify({'success': False, 'error': 'MySQL未连接'}), 500

    config_snapshot = data.get('config_snapshot') or {}
    if isinstance(config_snapshot, str):
        try:
            config_snapshot = json.loads(config_snapshot)
        except (json.JSONDecodeError, TypeError):
            return jsonify({'success': False, 'error': 'config_snapshot 格式错误'}), 400

    config_obj = config_snapshot.get('config') or {}
    if not isinstance(config_obj, dict):
        return jsonify({'success': False, 'error': 'config_snapshot 缺少 config 字段'}), 400

    generated = data.get('generated_content') or ''
    word_count = int(data.get('word_count') or 0) or (len(str(generated)) if generated else 0)
    scheme_description = (data.get('scheme_description') or '').strip() or None
    template_id = (data.get('template_id') or '') or None
    template_name = (data.get('template_name') or '') or None
    status = (data.get('generation_status') or 'complete').strip() or 'complete'
    task_id = (data.get('task_id') or f"jygz_{int(datetime.now().timestamp() * 1000)}_{uuid.uuid4().hex[:8]}")[:100]

    derived = _derive_jygz_selections_from_config(config_obj)
    env_map = config_obj.get('env_settings_map') or config_snapshot.get('editing_env_settings') or {}
    if isinstance(env_map, dict):
        environment_config = json.dumps(env_map, ensure_ascii=False, default=str)
    else:
        environment_config = None
    try:
        with Config.MYSQL_POOL.connection() as conn:
            _ensure_jygz_history_table(conn)
            with conn.cursor() as cursor:
                cursor.execute(
                    """
                    INSERT INTO jygz_history (
                        task_id, user_id, scheme_name, scheme_description,
                        selected_equipment, selected_environments, selected_projects,
                        selected_indicators, selected_methods, selected_files,
                        environment_config, template_id, template_name, config_snapshot,
                        generated_content, word_count, generation_status, completed_at
                    ) VALUES (
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
                    )
                    """,
                    (
                        task_id,
                        user_id,
                        scheme_name,
                        scheme_description,
                        json.dumps(derived['selected_equipment'], ensure_ascii=False),
                        json.dumps(derived['selected_environments'], ensure_ascii=False),
                        json.dumps(derived['selected_projects'], ensure_ascii=False),
                        json.dumps(derived['selected_indicators'], ensure_ascii=False),
                        json.dumps(derived['selected_methods'], ensure_ascii=False),
                        json.dumps(data.get('selected_files') or [], ensure_ascii=False),
                        environment_config,
                        template_id,
                        template_name,
                        json.dumps(config_snapshot, ensure_ascii=False, default=str),
                        generated,
                        word_count,
                        status,
                        datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                    ),
                )
            conn.commit()
    except Exception as e:
        logger.error(f'保存 jygz_history 失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500
    return jsonify({'success': True, 'data': {'task_id': task_id}})


@jygz_bp.route('/history/list', methods=['GET'])
def list_jygz_history():
    """当前用户的 JYGZ 历史方案列表（侧栏/管理）。"""
    user_id = (request.args.get('user_id') or 'admin').strip()
    limit = min(int(request.args.get('limit', 50)), 200)
    if not Config.MYSQL_POOL:
        return jsonify({'success': False, 'error': 'MySQL未连接'}), 500
    try:
        with Config.MYSQL_POOL.connection() as conn:
            _ensure_jygz_history_table(conn)
            with conn.cursor() as cursor:
                cursor.execute(
                    """
                    SELECT task_id, scheme_name, scheme_description, word_count, generation_status, created_at, updated_at
                    FROM jygz_history
                    WHERE user_id = %s
                    ORDER BY updated_at DESC
                    LIMIT %s
                    """,
                    (user_id, limit),
                )
                rows = cursor.fetchall()
        out = []
        for r in rows or []:
            out.append({
                'task_id': r[0],
                'scheme_name': r[1] or '未命名方案',
                'scheme_description': (r[2] or '')[:200],
                'word_count': r[3] or 0,
                'generation_status': r[4] or 'complete',
                'created_at': str(r[5]) if r[5] is not None else '',
                'updated_at': str(r[6]) if r[6] is not None else '',
            })
        return jsonify({'success': True, 'data': out})
    except Exception as e:
        logger.error(f'列表 jygz_history 失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@jygz_bp.route('/history/<path:task_id>', methods=['GET', 'DELETE'])
def get_jygz_history(task_id: str):
    """GET：按 task_id 取一条；DELETE：删除当前用户下的该条历史。"""
    user_id = (request.args.get('user_id') or 'admin').strip()
    if not Config.MYSQL_POOL:
        return jsonify({'success': False, 'error': 'MySQL未连接'}), 500

    if request.method == 'DELETE':
        deleted = 0
        try:
            with Config.MYSQL_POOL.connection() as conn:
                _ensure_jygz_history_table(conn)
                with conn.cursor() as cursor:
                    cursor.execute(
                        'DELETE FROM jygz_history WHERE task_id = %s AND user_id = %s',
                        (task_id, user_id),
                    )
                    deleted = cursor.rowcount
                conn.commit()
        except Exception as e:
            logger.error(f'删除 jygz_history 失败: {e}')
            return jsonify({'success': False, 'error': str(e)}), 500
        if deleted:
            return jsonify({'success': True, 'message': '已删除'})
        return jsonify({'success': False, 'error': '未找到该历史记录或无权删除'}), 404

    try:
        with Config.MYSQL_POOL.connection() as conn:
            _ensure_jygz_history_table(conn)
            with conn.cursor() as cursor:
                cursor.execute(
                    """
                    SELECT task_id, user_id, scheme_name, scheme_description, environment_config, template_id, template_name,
                           config_snapshot, generated_content, word_count, generation_status, created_at, updated_at, completed_at
                    FROM jygz_history
                    WHERE task_id = %s AND user_id = %s
                    """,
                    (task_id, user_id),
                )
                row = cursor.fetchone()
        if not row:
            return jsonify({'success': False, 'error': '未找到该历史记录'}), 404
    except Exception as e:
        logger.error(f'查询 jygz_history 失败: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500

    (t_task_id, t_user_id, scheme_n, scheme_d, env_c, tmpl_id, tmpl_name,
     snap_raw, generated_content, word_count, gen_status, created_at, updated_at, completed_at) = row
    if isinstance(snap_raw, str) and snap_raw.strip():
        try:
            snap = json.loads(snap_raw)
        except (json.JSONDecodeError, TypeError):
            snap = None
    else:
        snap = snap_raw if isinstance(snap_raw, dict) else None

    def _t(v):
        if v is None:
            return None
        if hasattr(v, 'isoformat'):
            s = v.isoformat()
            return s.replace('T', ' ')[:19] if 'T' in s else s[:19]
        return str(v)

    return jsonify({
        'success': True,
        'data': {
            'task_id': t_task_id,
            'user_id': t_user_id,
            'scheme_name': scheme_n,
            'scheme_description': scheme_d,
            'environment_config': env_c,
            'template_id': tmpl_id,
            'template_name': tmpl_name,
            'config_snapshot': snap,
            'generated_content': generated_content,
            'word_count': word_count,
            'generation_status': gen_status,
            'created_at': _t(created_at),
            'updated_at': _t(updated_at),
            'completed_at': _t(completed_at),
        },
    })


@jygz_bp.route('/config/examples', methods=['GET'])
def get_example_config():
    """返回示例配置（Matrice 350 RTK 无人机）"""
    example = {
        "设备名称": "DJI Matrice 350 RTK",
        "编号": "UAV-M350-2024-001",
        "用途": "工业巡检、测绘、应急救援等任务",
        "环境描述": "高温、低温、湿热、低气压、振动",
        "任务来源": "XX研究院2024年度无人机环境适应性试验任务",
        "术语列表": ["RTK", "GNSS", "IMU", "FPV", "MTOW"],
        "依据标准列表": [
            {"标识": "GJB 150A-2009", "名称": "军用装备实验室环境试验方法", "来源": "国军标"},
            {"标识": "GJB 899A-2009", "名称": "可靠性鉴定和验收试验", "来源": "国军标"},
            {"标识": "GB/T 2423", "名称": "电工电子产品环境试验", "来源": "国标"}
        ],
        "环境列表": [
            {
                "名称": "高温",
                "设备名称": "高低温试验箱",
                "关键参数": "温度范围：-40℃~+85℃，容积≥1m³",
                "依据标准": "GJB 150.3A-2009",
                "辅助仪器": ["温度记录仪", "红外测温仪"],
                "试验项目": [
                    {
                        "名称": "高温工作试验",
                        "准备_设备": "试验箱、电源、监控设备",
                        "准备_样品状态": "设备处于正常工作状态，电池充满",
                        "准备_安全防护": "佩戴防烫手套，确保通风",
                        "前检测_功能": "检查所有功能正常",
                        "前检测_外观": "检查外壳无裂纹、变形",
                        "前检测_性能": "记录初始性能参数",
                        "步骤_参数": [
                            {"名称": "温度", "数值": "+70℃", "依据": "GJB 150.3A"},
                            {"名称": "保持时间", "数值": "2小时", "依据": "GJB 150.3A"}
                        ],
                        "步骤_流程": "将样品放入试验箱，以≤2℃/min速率升温至+70℃，保持2小时，期间每30分钟检测一次功能",
                        "恢复_方法": "以≤2℃/min速率降温至室温",
                        "恢复_处理": "取出样品，在标准大气条件下恢复2小时",
                        "后检查_功能": "检查所有功能是否正常",
                        "后检查_结构": "检查结构是否变形、开裂",
                        "后检查_绝缘": "检查电气绝缘性能",
                        "后检查_失效判定": "任一功能失效或结构损坏即判定不合格"
                    }
                ]
            },
            {
                "名称": "低温",
                "设备名称": "高低温试验箱",
                "关键参数": "温度范围：-40℃~+85℃，容积≥1m³",
                "依据标准": "GJB 150.4A-2009",
                "辅助仪器": ["温度记录仪"],
                "试验项目": [
                    {
                        "名称": "低温工作试验",
                        "准备_设备": "试验箱、电源、预热设备",
                        "准备_样品状态": "设备处于正常工作状态",
                        "准备_安全防护": "佩戴防冻手套",
                        "前检测_功能": "检查所有功能正常",
                        "前检测_外观": "检查外壳无裂纹",
                        "前检测_性能": "记录初始性能参数",
                        "步骤_参数": [
                            {"名称": "温度", "数值": "-40℃", "依据": "GJB 150.4A"},
                            {"名称": "保持时间", "数值": "4小时", "依据": "GJB 150.4A"}
                        ],
                        "步骤_流程": "将样品放入试验箱，以≤1℃/min速率降温至-40℃，保持4小时，期间每30分钟检测一次功能",
                        "恢复_方法": "以≤1℃/min速率升温至室温",
                        "恢复_处理": "取出样品，在标准大气条件下恢复4小时",
                        "后检查_功能": "检查所有功能是否正常",
                        "后检查_结构": "检查材料是否脆化、开裂",
                        "后检查_绝缘": "检查电气绝缘性能",
                        "后检查_失效判定": "任一功能失效或结构损坏即判定不合格"
                    }
                ]
            }
        ],
        "边界性能指标": [
            {
                "名称": "最大飞行速度",
                "设备配置": "GPS定位系统、风速仪、秒表",
                "步骤": "在无风条件下，设置最大速度模式，记录水平飞行最大速度，重复3次取平均值",
                "合格判据": "最大速度≥15m/s"
            },
            {
                "名称": "最大续航时间",
                "设备配置": "计时器、电池监测系统",
                "步骤": "满电状态下，以悬停模式飞行至电池低电量报警，记录飞行时间",
                "合格判据": "悬停续航时间≥55分钟"
            },
            {
                "名称": "最大载荷",
                "设备配置": "电子秤、配重块",
                "步骤": "逐步增加配重，测试不同载荷下的飞行稳定性，记录最大安全载荷",
                "合格判据": "最大载荷≥2.7kg"
            }
        ]
    }
    
    return jsonify({'success': True, 'data': example})


# ==================== 辅助函数 ====================

def markdown_to_pdf(markdown_content: str, title: str) -> bytes:
    """将Markdown转换为PDF（与 JYLLM 共用 scheme_pdf_export，版式由 md2pdf-formal 统一）。"""
    try:
        from utils.scheme_pdf_export import export_markdown_to_pdf_bytes

        pdf_content, err = export_markdown_to_pdf_bytes(
            markdown_content, doc_title=title or "试验方案", doc_id=""
        )
        if pdf_content:
            return pdf_content
        logger.error("PDF生成失败: %s", err)
        return generate_simple_pdf(markdown_content, title)
    except Exception as e:
        logger.error(f"使用 scheme_pdf_export 失败: {e}")
        return generate_simple_pdf(markdown_content, title)


def generate_simple_pdf(content: str, title: str) -> bytes:
    """简单的PDF生成（使用reportlab）"""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import io
        
        # 尝试注册中文字体
        try:
            pdfmetrics.registerFont(TTFont('SimSun', 'C:/Windows/Fonts/simsun.ttc'))
            font_name = 'SimSun'
        except:
            font_name = 'Helvetica'
        
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4
        
        # 设置字体
        c.setFont(font_name, 16)
        c.drawString(100, height - 100, title)
        
        c.setFont(font_name, 12)
        y = height - 150
        
        for line in content.split('\n'):
            if y < 50:
                c.showPage()
                y = height - 50
                c.setFont(font_name, 12)
            
            # 简单的标题识别
            if line.startswith('# '):
                c.setFont(font_name, 16)
                c.drawString(50, y, line[2:])
                c.setFont(font_name, 12)
            elif line.startswith('## '):
                c.setFont(font_name, 14)
                c.drawString(50, y, line[3:])
                c.setFont(font_name, 12)
            elif line.startswith('### '):
                c.setFont(font_name, 13)
                c.drawString(50, y, line[4:])
                c.setFont(font_name, 12)
            else:
                c.drawString(50, y, line[:80])  # 限制每行长度
            
            y -= 20
        
        c.save()
        return buffer.getvalue()
    except Exception as e:
        logger.error(f"PDF生成失败: {e}")
        return b''


def get_default_template() -> str:
    """获取默认模板"""
    default_template_path = TEMPLATES_DIR / 'default_template.md'
    if default_template_path.exists():
        return default_template_path.read_text(encoding='utf-8')
    return get_builtin_default_template()


def get_builtin_default_template() -> str:
    """内置默认模板"""
    return '''# 【通用模板】环境及边界性能试验方案

## 1. 范围

### 1.1 标识
- **文档标识**：`{{编号}}`
- **文档标题**：`{{设备名称}}`环境及边界性能试验方案
- **术语和缩略语**：{{术语列表 | join('、')}}

### 1.2 试验目的
考核`{{设备名称}}`在`{{环境描述}}`下的环境适应性和边界性能，发现缺陷，提出改进和使用建议，为`{{用途}}`提供依据。

### 1.3 文档概述
本方案规定了`{{设备名称}}`的环境及边界性能试验的依据、环境与设备、内容与方法、中断恢复、结果评定等，是试验实施的依据。

## 2. 试验依据

### 2.1 任务来源
`{{任务来源}}`

### 2.2 依据文件
| 序号 | 文件标识 | 文件名称 | 来源 |
|------|----------|----------|------|
{% for std in 依据标准列表 %}
| {{ loop.index }} | {{ std.标识 }} | {{ std.名称 }} | {{ std.来源 }} |
{% endfor %}
{% if not 依据标准列表 %}
| 1 | 待补充 | 待补充 | 待补充 |
{% endif %}

## 3. 试验环境与设备

{% for env in 环境列表 %}
### 3.{{ loop.index }} {{ env.名称 }}试验设备

- **设备名称**：`{{ env.设备名称 }}`
- **关键参数**：{{ env.关键参数 }}
- **依据标准**：`{{ env.依据标准 }}`
- **辅助仪器**：{{ env.辅助仪器 | join('、') }}

{% if env.selected_subtitles %}
**分类选项**：
{% for subtitle in env.selected_subtitles %}
- {{ subtitle.letter }}) {{ subtitle.title }}：{{ subtitle.content }}
{% endfor %}
{% endif %}

{% endfor %}

## 4. 试验内容及方法

{% for env in 环境列表 %}
{% set env_idx = loop.index %}
### 4.{{ env_idx }} {{ env.名称 }}

{% for project in env.试验项目 %}
{% set proj_idx = loop.index %}
#### 4.{{ env_idx }}.{{ proj_idx }} {{ project.名称 }}

##### 4.{{ env_idx }}.{{ proj_idx }}.1 试验方法

**（1）试验准备**
- 设备与仪器：{{ project.准备_设备 }}
- 样品状态：{{ project.准备_样品状态 }}
- 安全防护：{{ project.准备_安全防护 }}

**（2）试验前检测**
- 功能自检：{{ project.前检测_功能 }}
- 外观与结构检查：{{ project.前检测_外观 }}
- 性能初测：{{ project.前检测_性能 }}

**（3）试验步骤**
- **参数设置**：

| 参数 | 数值 | 依据标准 |
|------|------|----------|
{% for param in project.步骤_参数 %}
| {{ param.名称 }} | {{ param.数值 }} | {{ param.依据 }} |
{% endfor %}

- **操作流程**：{{ project.步骤_流程 }}

**（4）恢复**
- 环境恢复方法：{{ project.恢复_方法 }}
- 样品处理：{{ project.恢复_处理 }}

**（5）试验后检查**
- 功能复检：{{ project.后检查_功能 }}
- 结构/密封分析：{{ project.后检查_结构 }}
- 绝缘/安全测试：{{ project.后检查_绝缘 }}
- **失效判定**：{{ project.后检查_失效判定 }}

{% endfor %}

{% endfor %}

{% set boundary_idx = 环境列表|length + 1 %}
### 4.{{ boundary_idx }} 边界性能考核

{% for metric in 边界性能指标 %}
{% set metric_idx = loop.index %}
#### 4.{{ boundary_idx }}.{{ metric_idx }} {{ metric.名称 }}

##### 4.{{ boundary_idx }}.{{ metric_idx }}.1 试验方法

**（1）试验设备与配置**：{{ metric.设备配置 }}

**（2）试验步骤**：{{ metric.步骤 }}

**（3）合格判据**：{{ metric.合格判据 }}

{% endfor %}

## 5. 试验中断与恢复

### 5.1 欠试验中断与恢复
试验条件低于允差下限时，从低于点重新达到规定条件后继续。

### 5.2 过试验中断与恢复
试验条件超过允差上限时，停止试验。若分析确认设备未损坏且无影响，可从超差点恢复继续；否则用新样机重做。

### 5.3 试验终止
- 正常终止：全部试验完成。
- 异常终止：样机损坏无法继续后续试验。

## 6. 试验结果评定

`{{设备名称}}`在各项环境试验前、后的外观/结构检查结果以及功能/性能检测结果满足本方案第4章中各项目规定的合格判据时，环境试验结果为合格。各项边界性能指标满足或优于指标要求时，边界性能试验结果为合格。

**总体结论**：待试验后判定。
'''


# ==================== 数据源API（复用JYLLM数据） ====================

def serialize_neo4j_value(value):
    from neo4j.time import DateTime, Date, Time
    if isinstance(value, (DateTime, Date, Time)):
        return value.isoformat()
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return value

@jygz_bp.route('/equipment-tree', methods=['GET'])
def get_equipment_tree():
    if not Config.NEO4J_DRIVER:
        return jsonify({"success": False, "error": "Neo4j未连接"})
    
    try:
        with Config.NEO4J_DRIVER.session() as session:
            result = session.run("""
                MATCH (et:EquipmentType)-[:HAS]->(e:Equipment)
                WHERE (et).Status IN $flgl_enabled_status AND (e).Status IN $flgl_enabled_status
                RETURN et, e
                ORDER BY et.Name, e.Name
            """, {**flgl_status_cypher_param()})
            
            tree = {}
            for record in result:
                et_data = {k: serialize_neo4j_value(v) for k, v in record["et"].items()}
                e_data = {k: serialize_neo4j_value(v) for k, v in record["e"].items()}
                
                et_type = et_data.get('Type', '未分类')
                et_id = et_data.get('Id', '')
                et_name = et_data.get('Name', et_id)
                
                eq_id = e_data.get('Id', '')
                eq_name = e_data.get('Name', e_data.get('name', '未命名'))
                eq_model = e_data.get('Model', e_data.get('model', ''))
                eq_type = e_data.get('Type', e_data.get('type', ''))
                
                if et_type not in tree:
                    tree[et_type] = {'label': et_type, 'types': {}}
                if et_id not in tree[et_type]['types']:
                    tree[et_type]['types'][et_id] = {'type': et_data, 'equipment': []}
                raw_data = {}
                for k, v in e_data.items():
                    raw_data[k] = serialize_neo4j_value(v)
                
                tree[et_type]['types'][et_id]['equipment'].append({
                    'id': eq_id,
                    'name': eq_name,
                    'model': eq_model,
                    'type': eq_type,
                    'raw': raw_data
                })
            
            result_tree = []
            for et_type, info in tree.items():
                types_list = []
                for et_id, et_info in info['types'].items():
                    types_list.append({
                        'title': et_info['type'].get('Name', et_id),
                        'key': f'et_{et_id}',
                        'children': [
                            {
                                'title': eq.get('name', eq.get('id', '')),
                                'key': eq.get('id', ''),
                                'isLeaf': True,
                                'model': eq.get('model', ''),
                                'type': eq.get('type', ''),
                                'data': eq.get('raw', {})
                            }
                            for eq in et_info['equipment']
                        ],
                    })
                result_tree.append({
                    'title': info['label'],
                    'key': f'type_{et_type}',
                    'children': types_list,
                })
            
        return jsonify({"success": True, "data": result_tree})
    except Exception as e:
        logger.error(f"获取设备树失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@jygz_bp.route('/environments', methods=['GET'])
def get_environments():
    if not Config.NEO4J_DRIVER:
        return jsonify({"success": False, "error": "Neo4j未连接"})
    
    try:
        with Config.NEO4J_DRIVER.session() as session:
            result = session.run("""
                MATCH (env:Environment)
                WHERE (env).Status IN $flgl_enabled_status
                RETURN env
                ORDER BY env.Name
            """, {**flgl_status_cypher_param()})
            
            environments = []
            for record in result:
                env_data = {k: serialize_neo4j_value(v) for k, v in record["env"].items()}
                
                env_id = env_data.get('Id', '')
                env_name = env_data.get('Name', '未命名')
                env_type = env_data.get('Type', '')
                env_class = env_data.get('Class', '')
                env_limit = env_data.get('Limit', '')
                env_request = env_data.get('Request', '')
                
                constraints_template = []
                requirements = []
                
                if env_limit:
                    try:
                        if isinstance(env_limit, str):
                            if env_limit.startswith('[') or env_limit.startswith('{'):
                                limit_data = json.loads(env_limit)
                                if isinstance(limit_data, list):
                                    for item in limit_data:
                                        if isinstance(item, dict):
                                            constraints_template.append({
                                                "key": item.get('key', ''),
                                                "label": item.get('label', ''),
                                                "unit": item.get('unit', ''),
                                                "type": item.get('type', 'value'),
                                                "placeholder": item.get('placeholder', '')
                                            })
                                        else:
                                            constraints_template.append({
                                                "key": f"limit_{len(constraints_template)}",
                                                "label": str(item),
                                                "unit": "",
                                                "type": "value",
                                                "placeholder": ""
                                            })
                                elif isinstance(limit_data, dict):
                                    for key, value in limit_data.items():
                                        constraints_template.append({
                                            "key": key,
                                            "label": value.get('label', key),
                                            "unit": value.get('unit', ''),
                                            "type": value.get('type', 'value'),
                                            "placeholder": value.get('placeholder', '')
                                        })
                            else:
                                for item in env_limit.split(','):
                                    item = item.strip()
                                    if item:
                                        constraints_template.append({
                                            "key": f"limit_{len(constraints_template)}",
                                            "label": item,
                                            "unit": "",
                                            "type": "value",
                                            "placeholder": ""
                                        })
                    except json.JSONDecodeError:
                        constraints_template.append({
                            "key": "limit_0",
                            "label": env_limit,
                            "unit": "",
                            "type": "value",
                            "placeholder": ""
                        })
                
                if env_request:
                    try:
                        if isinstance(env_request, str):
                            if env_request.startswith('[') or env_request.startswith('{'):
                                request_data = json.loads(env_request)
                                if isinstance(request_data, list):
                                    requirements = [str(item) for item in request_data]
                                elif isinstance(request_data, dict):
                                    requirements = list(request_data.keys())
                            else:
                                requirements = [item.strip() for item in env_request.split(',') if item.strip()]
                    except json.JSONDecodeError:
                        requirements = [env_request]
                
                if not constraints_template:
                    constraints_template = [
                        {"key": "default_param", "label": "参数", "unit": "", "type": "value", "placeholder": "请输入"}
                    ]
                
                if not requirements:
                    requirements = ["功能正常", "结构完整"]
                
                color_map = {
                    '温度': '#ef4444',
                    '湿度': '#10b981',
                    '压力': '#8b5cf6',
                    '力学': '#f59e0b',
                    '电磁': '#3b82f6',
                    '其他': '#6b7280'
                }
                
                icon_map = {
                    '温度': 'FireOutlined',
                    '湿度': 'CloudDownloadOutlined',
                    '压力': 'VerticalLeftOutlined',
                    '力学': 'AudioOutlined',
                    '电磁': 'SafetyOutlined',
                    '其他': 'ExperimentOutlined'
                }
                
                env_category = '其他'
                for key in color_map.keys():
                    if key in env_name or key in env_type:
                        env_category = key
                        break
                
                type_subtitles = []
                if env_type:
                    import re
                    pattern = r'([a-zA-Z])\)\s*(.+?)(?=\s*[a-zA-Z]\)|\s*$)'
                    matches = re.findall(pattern, env_type, re.DOTALL)
                    
                    if matches:
                        for letter, content in matches:
                            content = content.strip()
                            if ':' in content:
                                title, desc = content.split(':', 1)
                                title = title.strip()
                                desc = desc.strip()
                            else:
                                title = content[:20] + "..." if len(content) > 20 else content
                                desc = content
                            
                            type_subtitles.append({
                                "key": f"type_{letter.lower()}",
                                "title": title,
                                "content": desc,
                                "letter": letter.upper()
                            })
                    else:
                        type_subtitles.append({
                            "key": "type_main",
                            "title": env_type[:30] + "..." if len(env_type) > 30 else env_type,
                            "content": env_type,
                            "letter": "A"
                        })
                
                environments.append({
                    "id": env_id,
                    "name": env_name,
                    "category": env_type or env_category,
                    "icon": icon_map.get(env_category, 'ExperimentOutlined'),
                    "color": color_map.get(env_category, '#6b7280'),
                    "grades": ["轻度", "中等", "严酷"],
                    "default_grade": "中等",
                    "constraints_template": constraints_template,
                    "requirements": requirements,
                    "standard": env_data.get('Standard', ''),
                    "type_subtitles": type_subtitles,
                    "raw_data": env_data
                })
            
        return jsonify({"success": True, "data": environments})
    except Exception as e:
        logger.error(f"获取环境数据失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@jygz_bp.route('/projects', methods=['GET'])
def get_projects():
    """与 JQCX 一致：返回带 children/_env 的树，供配置表按真实环境列分组。"""
    from blueprints.GLCX.JQCX.routes import get_projects as jqcx_get_projects

    return jqcx_get_projects()


@jygz_bp.route('/indicators', methods=['GET'])
def get_indicators():
    if not Config.NEO4J_DRIVER:
        return jsonify({"success": False, "error": "Neo4j未连接"})
    
    equipment_id = request.args.get('equipment_id', '')
    logger.info(f"获取指标列表，设备ID: {equipment_id}")
    
    try:
        with Config.NEO4J_DRIVER.session() as session:
            if equipment_id:
                # 首先检查设备是否存在
                check_result = session.run("""
                    MATCH (e:Equipment {Id: $equipment_id})
                    WHERE (e).Status IN $flgl_enabled_status
                    RETURN e.Id as id, e.Name as Name
                """, {**flgl_status_cypher_param(), "equipment_id": equipment_id})
                equipment_record = check_result.single()
                if equipment_record:
                    logger.info(f"找到设备: {equipment_record['id']} - {equipment_record['Name']}")
                else:
                    logger.warning(f"未找到设备: {equipment_id}")
                    # 尝试其他可能的ID格式
                    check_result2 = session.run("""
                        MATCH (e:Equipment)
                        WHERE (e.Id = $equipment_id OR e.id = $equipment_id)
                          AND (e).Status IN $flgl_enabled_status
                        RETURN e.Id as id, e.Name as Name
                    """, {**flgl_status_cypher_param(), "equipment_id": equipment_id})
                    equipment_record2 = check_result2.single()
                    if equipment_record2:
                        logger.info(f"找到设备（备选查询）: {equipment_record2['id']} - {equipment_record2['Name']}")
                    else:
                        logger.warning(f"设备确实不存在: {equipment_id}")
                
                result = session.run("""
                    MATCH (e:Equipment {Id: $equipment_id})-[:CON]->(tp:TestProject)-[:CON]->(m:TestMethod)<-[:CON]-(i:Indicator)
                    WHERE (e).Status IN $flgl_enabled_status
                      AND (tp).Status IN $flgl_enabled_status
                      AND (m).Status IN $flgl_enabled_status
                      AND (i).Status IN $flgl_enabled_status
                    RETURN DISTINCT i.Id as id, i.Name as Name, i.description as description,
                           i.unit as unit, i.category as category
                    ORDER BY i.Name
                """, {**flgl_status_cypher_param(), "equipment_id": equipment_id})
            else:
                result = session.run("""
                    MATCH (i:Indicator)
                    WHERE (i).Status IN $flgl_enabled_status
                    RETURN i.Id as id, i.Name as Name, i.description as description,
                           i.unit as unit, i.category as category
                    ORDER BY i.Name
                """, {**flgl_status_cypher_param()})
            
            indicators = []
            for record in result:
                indicators.append({
                    "id": serialize_neo4j_value(record["id"]) or "",
                    "name": serialize_neo4j_value(record["Name"]) or "未命名",
                    "description": serialize_neo4j_value(record["description"]) or "",
                    "unit": serialize_neo4j_value(record["unit"]) or "",
                    "category": serialize_neo4j_value(record["category"]) or ""
                })
            
            logger.info(f"找到 {len(indicators)} 个指标")
        return jsonify({"success": True, "data": indicators})
    except Exception as e:
        logger.error(f"获取指标列表失败: {e}")
        import traceback
        logger.error(traceback.format_exc())
        return jsonify({"success": False, "error": str(e)}), 500


def _jygz_norm_id(v):
    if v is None:
        return ''
    return str(v).strip()


def _jygz_method_node_props(node):
    """从 Neo4j TestMethod 节点提取常用字段（小写键）。"""
    if node is None:
        return {}
    props = {}
    for k, v in dict(node).items():
        if k == 'labels':
            continue
        lk = k.lower()
        if lk in (
            'id', 'name', 'label', 'description', 'type', 'status', 'sy',
            'sytj', 'sygc', 'jgpd', 'synr',
        ):
            props[lk] = serialize_neo4j_value(v)
    return props


@jygz_bp.route('/methods', methods=['POST'])
def post_jygz_methods_batch():
    """
    与 JQCX POST /methods 一致：使用 GLCX methods_shared，
    返回 related_projects / related_indicators（含 rate）、max_rate、_relationType、files 等。
    """
    if not Config.NEO4J_DRIVER:
        return jsonify({"success": False, "error": "Neo4j未连接"}), 500

    try:
        from blueprints.GLCX.methods_shared import build_methods_by_project_and_indicators
        from blueprints.GLCX.JQCX.routes import _run_cypher, _get_node_props, _get_source_name

        data = request.get_json() or {}
        project_ids = [_jygz_norm_id(x) for x in (data.get('projectIds') or []) if _jygz_norm_id(x)]
        indicator_ids = [_jygz_norm_id(x) for x in (data.get('indicatorIds') or []) if _jygz_norm_id(x)]

        if not project_ids and not indicator_ids:
            return jsonify({"success": True, "data": []})

        nodes = build_methods_by_project_and_indicators(
            project_ids,
            indicator_ids,
            run_cypher=_run_cypher,
            get_node_props=_get_node_props,
            get_source_name=_get_source_name,
            files_include_source_id=False,
            populate_source_name=False,
        )
        return jsonify({"success": True, "data": nodes})
    except Exception as e:
        logger.error(f"JYGZ POST /methods 失败: {e}", exc_info=True)
        return jsonify({"success": False, "error": str(e)}), 500


@jygz_bp.route('/methods', methods=['GET'])
def get_methods():
    if not Config.NEO4J_DRIVER:
        return jsonify({"success": False, "error": "Neo4j未连接"})
    
    project_id = request.args.get('project_id', '')
    indicator_id = request.args.get('indicator_id', '')
    
    try:
        with Config.NEO4J_DRIVER.session() as session:
            if project_id or indicator_id:
                project_ids = [project_id] if project_id else []
                indicator_ids = [indicator_id] if indicator_id else []
                
                project_methods = {}
                indicator_methods = {}
                
                if project_ids:
                    result = session.run("""
                        MATCH (tp:TestProject)-[:CON]->(m:TestMethod)
                        WHERE tp.Id IN $project_ids
                          AND (tp).Status IN $flgl_enabled_status
                          AND (m).Status IN $flgl_enabled_status
                        OPTIONAL MATCH (m)-[:BASED_ON]->(f:Document)
                        RETURN DISTINCT m.Id as id, m.Name as Name, m.Description as description,
                               m.Type as type, m.Label as label, m.Sy as sy, tp.Name as project_Name,
                               m.SYTJ as SYTJ, m.SYGC as SYGC, m.JGPD as JGPD, m.SYNR as SYNR,
                               collect(DISTINCT {id: f.Id, name: f.Name}) as files
                    """, {**flgl_status_cypher_param(), "project_ids": project_ids})
                    
                    for record in result:
                        logger.info(f"Method record: {dict(record)}")
                        method_id = serialize_neo4j_value(record["id"])
                        files_data = []
                        for f in record.get("files", []):
                            if f and f.get("id"):
                                files_data.append({
                                    "id": serialize_neo4j_value(f.get("id")),
                                    "name": serialize_neo4j_value(f.get("name")) or ""
                                })
                        if method_id not in project_methods:
                            project_methods[method_id] = {
                                "id": method_id,
                                "name": serialize_neo4j_value(record["Name"]) or "未命名",
                                "description": serialize_neo4j_value(record["description"]) or "",
                                "type": serialize_neo4j_value(record["type"]) or "",
                                "label": serialize_neo4j_value(record["label"]) or "",
                                "sy": serialize_neo4j_value(record["sy"]) or "",
                                "SYTJ": serialize_neo4j_value(record.get("SYTJ")) or "",
                                "SYGC": serialize_neo4j_value(record.get("SYGC")) or "",
                                "JGPD": serialize_neo4j_value(record.get("JGPD")) or "",
                                "SYNR": serialize_neo4j_value(record.get("SYNR")) or "",
                                "project_sources": [],
                                "indicator_sources": [],
                                "files": files_data
                            }
                        project_name = serialize_neo4j_value(record["project_Name"])
                        if project_name and project_name not in project_methods[method_id]["project_sources"]:
                            project_methods[method_id]["project_sources"].append(project_name)
                
                if indicator_ids:
                    result = session.run("""
                        MATCH (i:Indicator)-[:CON]->(m:TestMethod)
                        WHERE i.Id IN $indicator_ids
                          AND (i).Status IN $flgl_enabled_status
                          AND (m).Status IN $flgl_enabled_status
                        OPTIONAL MATCH (m)-[:BASED_ON]->(f:Document)
                        RETURN DISTINCT m.Id as id, m.Name as Name, m.Description as description,
                               m.Type as type, m.Label as label, m.Sy as sy, i.Name as indicator_Name,
                               m.SYTJ as SYTJ, m.SYGC as SYGC, m.JGPD as JGPD, m.SYNR as SYNR,
                               collect(DISTINCT {id: f.Id, name: f.Name}) as files
                    """, {**flgl_status_cypher_param(), "indicator_ids": indicator_ids})
                    
                    for record in result:
                        method_id = serialize_neo4j_value(record["id"])
                        files_data = []
                        for f in record.get("files", []):
                            if f and f.get("id"):
                                files_data.append({
                                    "id": serialize_neo4j_value(f.get("id")),
                                    "name": serialize_neo4j_value(f.get("name")) or ""
                                })
                        if method_id not in indicator_methods:
                            indicator_methods[method_id] = {
                                "id": method_id,
                                "name": serialize_neo4j_value(record["Name"]) or "未命名",
                                "description": serialize_neo4j_value(record["description"]) or "",
                                "type": serialize_neo4j_value(record["type"]) or "",
                                "label": serialize_neo4j_value(record["label"]) or "",
                                "sy": serialize_neo4j_value(record["sy"]) or "",
                                "SYTJ": serialize_neo4j_value(record.get("SYTJ")) or "",
                                "SYGC": serialize_neo4j_value(record.get("SYGC")) or "",
                                "JGPD": serialize_neo4j_value(record.get("JGPD")) or "",
                                "SYNR": serialize_neo4j_value(record.get("SYNR")) or "",
                                "project_sources": [],
                                "indicator_sources": [],
                                "files": files_data
                            }
                        indicator_name = serialize_neo4j_value(record["indicator_Name"])
                        if indicator_name and indicator_name not in indicator_methods[method_id]["indicator_sources"]:
                            indicator_methods[method_id]["indicator_sources"].append(indicator_name)
                
                methods_dict = {**project_methods, **indicator_methods}
                for method_id, method_data in methods_dict.items():
                    if method_id in project_methods and method_id in indicator_methods:
                        method_data["project_sources"] = project_methods[method_id]["project_sources"]
                        method_data["indicator_sources"] = indicator_methods[method_id]["indicator_sources"]
                        all_files = project_methods[method_id].get("files", []) + indicator_methods[method_id].get("files", [])
                        seen_ids = set()
                        unique_files = []
                        for f in all_files:
                            if f["id"] not in seen_ids:
                                seen_ids.add(f["id"])
                                unique_files.append(f)
                        method_data["files"] = unique_files
                        if not method_data.get("sy") and project_methods[method_id].get("sy"):
                            method_data["sy"] = project_methods[method_id]["sy"]
                        if not method_data.get("sy") and indicator_methods[method_id].get("sy"):
                            method_data["sy"] = indicator_methods[method_id]["sy"]
                        for fld in ("SYTJ", "SYGC", "JGPD", "SYNR"):
                            if not method_data.get(fld) and project_methods[method_id].get(fld):
                                method_data[fld] = project_methods[method_id][fld]
                            if not method_data.get(fld) and indicator_methods[method_id].get(fld):
                                method_data[fld] = indicator_methods[method_id][fld]
                
                methods = []
                for method_data in methods_dict.values():
                    project_sources = method_data["project_sources"]
                    indicator_sources = method_data["indicator_sources"]
                    
                    if project_sources and indicator_sources:
                        relation_type = "both"
                    elif project_sources:
                        relation_type = "project"
                    elif indicator_sources:
                        relation_type = "indicator"
                    else:
                        relation_type = "none"
                    
                    methods.append({
                        "id": method_data["id"],
                        "name": method_data["name"],
                        "description": method_data["description"],
                        "type": method_data["type"],
                        "label": method_data["label"],
                        "sy": method_data.get("sy", ""),
                        "SYTJ": method_data.get("SYTJ", ""),
                        "SYGC": method_data.get("SYGC", ""),
                        "JGPD": method_data.get("JGPD", ""),
                        "SYNR": method_data.get("SYNR", ""),
                        "relationType": relation_type,
                        "sourceProject": ", ".join(project_sources) if project_sources else "",
                        "sourceIndicator": ", ".join(indicator_sources) if indicator_sources else "",
                        "files": method_data.get("files", [])
                    })
            else:
                result = session.run("""
                    MATCH (m:TestMethod)
                    WHERE (m).Status IN $flgl_enabled_status
                    OPTIONAL MATCH (m)-[:BASED_ON]->(f:Document)
                    RETURN m.Id as id, m.Name as Name, m.Description as description,
                           m.Type as type, m.Label as label, m.Sy as sy,
                           m.SYTJ as SYTJ, m.SYGC as SYGC, m.JGPD as JGPD, m.SYNR as SYNR,
                           collect(DISTINCT {id: f.Id, name: f.Name}) as files
                    ORDER BY m.Name
                """, {**flgl_status_cypher_param()})
                
                methods = []
                for record in result:
                    files_data = []
                    for f in record.get("files", []):
                        if f and f.get("id"):
                            files_data.append({
                                "id": serialize_neo4j_value(f.get("id")),
                                "name": serialize_neo4j_value(f.get("name")) or ""
                            })
                    methods.append({
                        "id": serialize_neo4j_value(record["id"]) or "",
                        "name": serialize_neo4j_value(record["Name"]) or "未命名",
                        "description": serialize_neo4j_value(record["description"]) or "",
                        "type": serialize_neo4j_value(record["type"]) or "",
                        "label": serialize_neo4j_value(record["label"]) or "",
                        "sy": serialize_neo4j_value(record["sy"]) or "",
                        "SYTJ": serialize_neo4j_value(record.get("SYTJ")) or "",
                        "SYGC": serialize_neo4j_value(record.get("SYGC")) or "",
                        "JGPD": serialize_neo4j_value(record.get("JGPD")) or "",
                        "SYNR": serialize_neo4j_value(record.get("SYNR")) or "",
                        "relationType": "none",
                        "sourceProject": "",
                        "sourceIndicator": "",
                        "files": files_data
                    })
            
        return jsonify({"success": True, "data": methods})
    except Exception as e:
        logger.error(f"获取方法列表失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@jygz_bp.route('/boundary-indicators', methods=['GET'])
def get_boundary_indicators():
    indicators = [
        {"id": "max_speed", "name": "最大飞行速度", "unit": "m/s", "default_value": "15", "typical_method": "速度测试标准方法"},
        {"id": "endurance", "name": "最大续航时间", "unit": "分钟", "default_value": "55", "typical_method": "续航测试标准方法"},
        {"id": "max_payload", "name": "最大载荷", "unit": "kg", "default_value": "2.7", "typical_method": "载荷测试标准方法"},
        {"id": "wind_resistance", "name": "抗风等级", "unit": "级", "default_value": "6", "typical_method": "抗风测试标准方法"},
        {"id": "max_altitude", "name": "最大飞行高度", "unit": "m", "default_value": "5000", "typical_method": "高度测试标准方法"},
        {"id": "control_accuracy", "name": "控制精度", "unit": "m", "default_value": "0.1", "typical_method": "精度测试标准方法"},
        {"id": "response_time", "name": "响应时间", "unit": "ms", "default_value": "100", "typical_method": "响应测试标准方法"},
        {"id": "protection_level", "name": "防护等级", "unit": "IP", "default_value": "54", "typical_method": "防护等级测试方法"}
    ]
    
    return jsonify({"success": True, "data": indicators})


@jygz_bp.route('/recommend-environment', methods=['POST'])
def recommend_environment():
    if not Config.OPENAI_CLIENT:
        return jsonify({"success": False, "error": "LLM客户端未初始化"})
    
    data = request.get_json()
    equipment_info = data.get('equipment', {})
    environment_info = data.get('environment', {})
    
    equipment_name = equipment_info.get('name', '')
    equipment_type = equipment_info.get('type', '')
    equipment_model = equipment_info.get('model', '')
    env_name = environment_info.get('name', '')
    env_type = environment_info.get('type', '')
    
    prompt = f"""根据以下设备和环境信息，推荐合适的环境试验设置参数。

设备信息：
- 设备名称：{equipment_name}
- 设备类型：{equipment_type}
- 设备型号：{equipment_model}

环境类型：{env_name}
环境分类：{env_type}

请提供以下推荐设置：
1. 分级（轻度/中等/严酷）
2. 约束条件（如温度范围、湿度、持续时间等）
3. 要求（如功能正常、结构完整等）

请以JSON格式返回，包含以下字段：
- grade: 推荐的分级
- constraints: 约束条件对象，包含key-value对
- requirements: 要求数组
- reason: 推荐理由

示例格式：
{{
    "grade": "中等",
    "constraints": {{
        "temp_max": "+70",
        "temp_min": "+40",
        "duration": "2",
        "ramp_rate": "≤2"
    }},
    "requirements": ["功能正常", "结构完整", "绝缘良好"],
    "reason": "根据设备类型和环境要求，推荐中等严酷等级..."
}}"""
    
    try:
        response = Config.OPENAI_CLIENT.chat.completions.create(
            model=Config.OPENAI_MODEL,
            messages=[
                {"role": "system", "content": "你是一个专业的环境试验专家，能够根据设备类型和环境要求推荐合适的试验设置参数。"},
                {"role": "user", "content": prompt}
            ],
            temperature=0.3,
            max_tokens=1000
        )
        
        content = response.choices[0].message.content
        try:
            recommendation = json.loads(content)
            return jsonify({"success": True, "data": recommendation})
        except json.JSONDecodeError:
            return jsonify({"success": True, "data": {"raw_response": content}})
            
    except Exception as e:
        logger.error(f"LLM推荐失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


"""
以下解析方法已废弃，保留供参考
废弃原因：前端使用react-katex渲染LaTeX，HTML表格用dangerouslySetInnerHTML直接渲染
后端不再做格式化，只返回数据库原始字符串
"""

@jygz_bp.route('/environment-settings', methods=['POST'])
def get_environment_settings():
    data = request.get_json()
    if not data:
        return jsonify({"success": False, "error": "请求体为空"}), 400
    env_names = data.get('env_names', [])
    if not env_names:
        return jsonify({"success": False, "error": "环境名称列表不能为空"}), 400
    if not Config.NEO4J_DRIVER:
        return jsonify({"success": False, "error": "Neo4j未连接"}), 500
    try:
        with Config.NEO4J_DRIVER.session() as session:
            result = session.run("""
                MATCH (env:Environment)
                WHERE env.Name IN $env_names
                  AND (env).Status IN $flgl_enabled_status
                RETURN env.Name as name, 
                       env.Type as type, env.Class as class,
                       env.Limit as limit, env.Request as request,
                       env.Combination as combination
            """, {**flgl_status_cypher_param(), "env_names": env_names})
            settings_list = []
            for record in result:
                env_name = serialize_neo4j_value(record.get("name")) or "未命名"
                settings_list.append({
                    "env_name": env_name,
                    "type": serialize_neo4j_value(record.get("type")) or "",
                    "class": serialize_neo4j_value(record.get("class")) or "",
                    "limit": serialize_neo4j_value(record.get("limit")) or "",
                    "request": serialize_neo4j_value(record.get("request")) or "",
                    "combination": serialize_neo4j_value(record.get("combination")) or "",
                })
        return jsonify({"success": True, "data": settings_list})
    except Exception as e:
        logger.error(f"获取环境设置失败: {e}")
        import traceback
        logger.error(traceback.format_exc())
        return jsonify({"success": False, "error": str(e)}), 500
