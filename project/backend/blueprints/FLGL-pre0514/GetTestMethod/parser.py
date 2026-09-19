import os
import re
import logging
from collections import Counter
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    from docx import Document
except ImportError:
    Document = None


PAGE_MARKER_RE = re.compile(r"^\s*=+\s*Page\s+(\d+)\s*=+\s*$", re.IGNORECASE)
DECIMAL_HEADING_RE = re.compile(r"^\s*(\d+(?:\.\d+){0,5})[)\.、\s]+(.+?)\s*$")
CHAPTER_HEADING_RE = re.compile(r"^\s*第[一二三四五六七八九十百零0-9]+[章节部分篇]\s*(.+?)?\s*$")
APPENDIX_HEADING_RE = re.compile(r"^\s*附录\s*[A-Za-z0-9一二三四五六七八九十]?\s*(.*)$")
COMMON_HEADING_PREFIXES = (
    "范围", "规范性引用文件", "引证文件", "引用文件", "术语和定义", "术语", "定义", "缩略语",
    "总则", "通用要求", "目的和应用", "目的", "应用", "限制", "裁裁指南", "剪裁指南",
    "试验条件", "试验程序", "试验方法", "试验步骤", "试验过程", "试验要求", "信息要求",
    "试验中断", "中断", "恢复", "结果评定", "结果分析", "判定规则", "检验规则",
    "高温", "低温", "湿热", "振动", "冲击", "跌落", "贮存", "工作",
)
DEFAULT_SECTION_TITLE = "全文"

WATERMARK_HINTS = (
    "weisstech", "gjb标准", "中国可靠性网", "kekaoxing", "http://", "https://",
    "confidential", "internal use", "draft", "样张", "机密", "内部资料",
)
STANDARD_SIGNAL_WORDS = (
    "范围", "术语", "定义", "目的", "应用", "限制", "试验", "方法", "过程",
    "结果", "分析", "信息要求", "试验要求", "确定试验条件", "工作考虑", "表",
)


def _normalize_spaces(text: str) -> str:
    return re.sub(r"[ \t]+", " ", (text or "").replace("\r", "")).strip()


def _strip_table_tags(text: str) -> str:
    return re.sub(r"<[^>]+>", "", text or "").strip()


def _contains_chinese(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text or ""))


def _rebuild_full_text_from_pages(pages: List[Dict[str, Any]]) -> str:
    parts: List[str] = []
    for page in pages:
        text = (page.get("text") or "").strip()
        if text:
            parts.append(f"===== Page {page.get('page_number')} =====\n{text}")
    return "\n\n".join(parts).strip()


def _normalize_line_for_compare(line: str) -> str:
    s = _normalize_spaces(line).lower()
    s = re.sub(r"\s+", "", s)
    return s


def _is_probable_noise_line(line: str) -> bool:
    raw = _normalize_spaces(line)
    if not raw:
        return False
    lowered = raw.lower()
    if any(h in lowered for h in WATERMARK_HINTS):
        return True
    # 单独页码、罗马数字页码、过短的英文水印
    if re.fullmatch(r"[ivxlcdmIVXLCDM0-9]+", raw):
        return True
    if len(raw) <= 20 and re.fullmatch(r"[A-Za-z0-9_\-./: ]+", raw):
        return True
    return False


def _collect_repeated_noise_lines(pages: List[Dict[str, Any]], min_page_ratio: float = 0.6) -> set:
    total_pages = max(len(pages), 1)
    threshold = max(2, int(total_pages * min_page_ratio + 0.999))
    counter: Counter = Counter()
    original_map: Dict[str, str] = {}

    for page in pages:
        seen_once = set()
        for line in (page.get("text") or "").splitlines():
            key = _normalize_line_for_compare(line)
            if not key or key in seen_once:
                continue
            seen_once.add(key)
            counter[key] += 1
            original_map.setdefault(key, line)

    repeated = set()
    for key, cnt in counter.items():
        line = original_map.get(key, "")
        if cnt >= threshold and _is_probable_noise_line(line):
            repeated.add(key)
    return repeated


def _clean_pages_noise(pages: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not pages:
        return []

    repeated_noise = _collect_repeated_noise_lines(pages)
    cleaned_pages: List[Dict[str, Any]] = []

    for page in pages:
        kept_lines: List[str] = []
        for line in (page.get("text") or "").splitlines():
            key = _normalize_line_for_compare(line)
            if key in repeated_noise:
                continue
            kept_lines.append(line)

        cleaned_pages.append({
            **page,
            "text": "\n".join(kept_lines).strip(),
        })

    return cleaned_pages


def _pdf_text_quality(data: Dict[str, Any]) -> Dict[str, Any]:
    pages = list(data.get("pages", []) or [])
    full_text = (data.get("full_text") or "").strip()
    nonempty_pages = [p for p in pages if (p.get("text") or "").strip()]
    page_texts = [_normalize_spaces(p.get("text") or "") for p in nonempty_pages]
    unique_ratio = len(set(page_texts)) / max(len(page_texts), 1)

    chinese_chars = len(re.findall(r"[\u4e00-\u9fff]", full_text))
    keyword_hits = sum(1 for kw in STANDARD_SIGNAL_WORDS if kw in full_text)

    meaningful_pages = 0
    noise_pages = 0
    for t in page_texts:
        zh_count = len(re.findall(r"[\u4e00-\u9fff]", t))
        if zh_count >= 20 or any(kw in t for kw in STANDARD_SIGNAL_WORDS):
            meaningful_pages += 1
        if _is_probable_noise_line(t) or all(_is_probable_noise_line(x) for x in t.splitlines() if x.strip()):
            noise_pages += 1

    return {
        "page_count": len(pages),
        "nonempty_page_count": len(nonempty_pages),
        "full_text_len": len(full_text),
        "chinese_chars": chinese_chars,
        "keyword_hits": keyword_hits,
        "meaningful_pages": meaningful_pages,
        "unique_ratio": unique_ratio,
        "noise_pages": noise_pages,
    }


def _should_force_ocr(fitz_data: Dict[str, Any]) -> bool:
    q = _pdf_text_quality(fitz_data)

    if q["nonempty_page_count"] == 0:
        return True
    if q["full_text_len"] < 800:
        return True
    if q["chinese_chars"] < 120:
        return True
    if q["meaningful_pages"] < 2:
        return True
    if q["unique_ratio"] < 0.5:
        return True
    if q["noise_pages"] >= max(3, q["page_count"] // 2) and q["keyword_hits"] == 0:
        return True
    return False


def _is_heading_candidate(line: str) -> bool:
    s = _normalize_spaces(line)
    if not s:
        return False
    if len(s) > 80:
        return False
    if DECIMAL_HEADING_RE.match(s):
        return True
    if CHAPTER_HEADING_RE.match(s):
        return True
    if APPENDIX_HEADING_RE.match(s):
        return True
    if any(s.startswith(prefix) for prefix in COMMON_HEADING_PREFIXES):
        return True
    if re.match(r"^[（(]?[0-9一二三四五六七八九十]+[)）][、\s]?", s):
        return True
    return False


def _heading_level(line: str) -> int:
    s = _normalize_spaces(line)
    m = DECIMAL_HEADING_RE.match(s)
    if m:
        return len(m.group(1).split("."))
    if CHAPTER_HEADING_RE.match(s):
        return 1
    if APPENDIX_HEADING_RE.match(s):
        return 1
    return 2


def _read_txt(file_path: str) -> Dict[str, Any]:
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()
    return {
        "mode": "txt",
        "full_text": text.strip(),
        "pages": [{"page_number": 1, "text": text.strip(), "tables_html": [], "figures": []}],
        "all_tables": [],
        "ocr_used": False,
    }


def _read_docx(file_path: str) -> Dict[str, Any]:
    if Document is None:
        raise RuntimeError("缺少 python-docx，请先安装：pip install python-docx")

    doc = Document(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]

    tables_html: List[str] = []
    table_texts: List[str] = []
    for table in doc.tables:
        rows = []
        html_rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
                html_cells = "".join([f"<td>{cell}</td>" for cell in cells])
                html_rows.append(f"<tr>{html_cells}</tr>")
        if rows:
            table_texts.append("\n".join(rows))
            tables_html.append("<table>" + "".join(html_rows) + "</table>")

    parts: List[str] = []
    if paragraphs:
        parts.append("\n\n".join(paragraphs))
    if table_texts:
        parts.append("【表格文本】\n" + "\n\n".join(table_texts))

    full_text = "\n\n".join(parts).strip()
    page = {
        "page_number": 1,
        "text": full_text,
        "tables_html": tables_html,
        "figures": [],
    }
    return {
        "mode": "docx",
        "full_text": full_text,
        "pages": [page],
        "all_tables": [{"page": 1, "html": html} for html in tables_html],
        "ocr_used": False,
    }


def _read_pdf_with_fitz(file_path: str) -> Dict[str, Any]:
    if fitz is None:
        raise RuntimeError("缺少 PyMuPDF，请先安装：pip install pymupdf")

    doc = fitz.open(file_path)
    pages: List[Dict[str, Any]] = []
    all_text_parts: List[str] = []
    for i in range(len(doc)):
        page = doc[i]
        text = (page.get_text("text") or "").strip()
        pages.append({
            "page_number": i + 1,
            "text": text,
            "tables_html": [],
            "figures": [],
        })
        if text:
            all_text_parts.append(f"===== Page {i + 1} =====\n{text}")
    doc.close()

    return {
        "mode": "pdf_fitz",
        "full_text": "\n\n".join(all_text_parts).strip(),
        "pages": pages,
        "all_tables": [],
        "ocr_used": False,
    }


def _read_pdf_with_ocr(file_path: str, output_dir: str) -> Dict[str, Any]:
    # 延迟导入：避免仅走 OCR 缓存或 LLM 调试时加载 PaddleOCR/imgaug（与 numpy 2.x 常冲突）
    from pdf_ocr_extractor import OCRExtractorConfig, PDFOCRExtractor

    config = OCRExtractorConfig(
        pdf_image_dpi=300,
        ocr_use_gpu=False,
        output_dir=output_dir,
    )
    extractor = PDFOCRExtractor(config)
    result = extractor.extract(file_path)

    pages = []
    for page in result.get("pages", []):
        pages.append({
            "page_number": page.get("page_number"),
            "text": (page.get("text") or "").strip(),
            "tables_html": list(page.get("tables", []) or []),
            "figures": list(page.get("figures", []) or []),
        })

    pages = _clean_pages_noise(pages)
    full_text = _rebuild_full_text_from_pages(pages)

    return {
        "mode": "pdf_ocr",
        "full_text": full_text,
        "pages": pages,
        "all_tables": list(result.get("all_tables", []) or []),
        "ocr_used": True,
        "ocr_result": result,
    }


def _ocr_stem_for_pdf(pdf_path: str) -> str:
    """与 pdf_ocr_extractor.PDFOCRExtractor.extract 中目录名规则一致: ocr_{base_name}。"""
    pdf_name = os.path.basename(pdf_path)
    return pdf_name.replace(".pdf", "").replace(".PDF", "").replace("+", "_")


def ocr_cache_extracted_text_path(pdf_path: str, ocr_output_dir: str) -> str:
    """如 backend/output_json/ocr_tmp/ocr_GBT38924.3-2020/extracted_text.txt"""
    base = _ocr_stem_for_pdf(pdf_path)
    return os.path.join(ocr_output_dir, f"ocr_{base}", "extracted_text.txt")


def _parsed_from_cached_extracted_text(full_text: str) -> Optional[Dict[str, Any]]:
    """
    从历史上保存的 extracted_text.txt 内容还原与 pdf_ocr 兼容的页结构，便于跳过重跑 Paddle OCR。
    """
    s = (full_text or "").strip()
    if not s:
        return None
    page_start = re.compile(r"^\s*=+\s*Page\s+(\d+)\s*=+\s*$", re.IGNORECASE | re.MULTILINE)
    matches = list(page_start.finditer(s))
    if not matches:
        pages: List[Dict[str, Any]] = [
            {"page_number": 1, "text": s, "tables_html": [], "figures": []}
        ]
        pages = _clean_pages_noise(pages)
        return {
            "mode": "pdf_ocr",
            "full_text": _rebuild_full_text_from_pages(pages),
            "pages": pages,
            "all_tables": [],
            "ocr_used": True,
            "from_ocr_cache": True,
        }
    pages = []
    for i, m in enumerate(matches):
        pno = int(m.group(1))
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(s)
        chunk = s[start:end].strip()
        if chunk:
            pages.append(
                {
                    "page_number": pno,
                    "text": chunk,
                    "tables_html": [],
                    "figures": [],
                }
            )
    if not pages:
        return None
    pages = _clean_pages_noise(pages)
    return {
        "mode": "pdf_ocr",
        "full_text": _rebuild_full_text_from_pages(pages),
        "pages": pages,
        "all_tables": [],
        "ocr_used": True,
        "from_ocr_cache": True,
    }


def _try_load_parsed_from_ocr_cache(cache_path: str) -> Optional[Dict[str, Any]]:
    if not (cache_path and os.path.isfile(cache_path) and os.path.getsize(cache_path) > 0):
        return None
    try:
        with open(cache_path, "r", encoding="utf-8") as f:
            raw = f.read()
    except OSError:
        return None
    return _parsed_from_cached_extracted_text(raw)


def _ocr_failure_network_hint(ocr_error: Optional[Exception]) -> str:
    if ocr_error is None:
        return ""
    msg = str(ocr_error).lower()
    keys = (
        "getaddrinfo",
        "failed to resolve",
        "name resolution",
        "connection",
        "max retries",
        "timed out",
        "timeout",
        "bcebos",
        "443",
    )
    if not any(k in msg for k in keys):
        return ""
    return (
        " 离线/DNS 无法访问模型下载地址时：请在可联网环境预先初始化 PaddleOCR（自动下载到用户目录 "
        "`.paddleocr/whl/`），将整个 `.paddleocr` 拷贝到本机相同路径；或改用已有 OCR 缓存 "
        "`backend/output_json/ocr_tmp/ocr_<pdf主文件名>/extracted_text.txt` 后重试。"
    )


def _choose_pdf_result(
    fitz_data: Dict[str, Any],
    ocr_data: Optional[Dict[str, Any]],
    fitz_bad: bool,
    ocr_error: Optional[Exception],
) -> Dict[str, Any]:
    if ocr_data is None:
        if fitz_bad:
            detail = f"；OCR失败：{ocr_error}" if ocr_error else ""
            hint = _ocr_failure_network_hint(ocr_error)
            raise RuntimeError(
                "PDF文本层提取质量过低，且无法获得可用的OCR结果。"
                "当前继续运行只会得到页眉/水印等无效内容，请先检查PaddleOCR环境或PDF可读性"
                f"{detail}"
                f"{hint}"
            )
        return fitz_data

    fitz_q = _pdf_text_quality(fitz_data)
    ocr_q = _pdf_text_quality(ocr_data)

    # fitz明显差时，优先用OCR，只要OCR有基本可用内容
    if fitz_bad and (ocr_q["meaningful_pages"] >= 2 or ocr_q["chinese_chars"] >= 120):
        chosen = ocr_data
    else:
        # 一般情况下，比较质量而不是纯长度
        choose_ocr = False
        if ocr_q["meaningful_pages"] > fitz_q["meaningful_pages"]:
            choose_ocr = True
        elif ocr_q["keyword_hits"] > fitz_q["keyword_hits"]:
            choose_ocr = True
        elif ocr_q["chinese_chars"] > max(fitz_q["chinese_chars"] * 1.2, fitz_q["chinese_chars"] + 100):
            choose_ocr = True
        elif len(ocr_data.get("all_tables", []) or []) > 0 and fitz_q["keyword_hits"] == 0:
            choose_ocr = True

        chosen = ocr_data if choose_ocr else fitz_data

    # 若最终仍选fitz，也尽量把OCR表格挂上去
    if chosen is fitz_data and ocr_data is not None:
        ocr_pages_by_num = {p["page_number"]: p for p in ocr_data.get("pages", [])}
        enriched_pages = []
        for page in chosen.get("pages", []):
            current = dict(page)
            ocr_page = ocr_pages_by_num.get(page.get("page_number"))
            if ocr_page:
                current["tables_html"] = list(ocr_page.get("tables_html", []) or [])
                current["figures"] = list(ocr_page.get("figures", []) or [])
            enriched_pages.append(current)
        chosen = dict(chosen)
        chosen["pages"] = enriched_pages
        chosen["all_tables"] = list(ocr_data.get("all_tables", []) or [])
        chosen["full_text"] = _rebuild_full_text_from_pages(enriched_pages)

    chosen = dict(chosen)
    chosen["pages"] = _clean_pages_noise(list(chosen.get("pages", []) or []))
    chosen["full_text"] = _rebuild_full_text_from_pages(chosen["pages"])
    return chosen


def _build_page_table_map(pages: List[Dict[str, Any]]) -> Dict[int, List[str]]:
    result: Dict[int, List[str]] = {}
    for page in pages:
        page_no = int(page.get("page_number", 0) or 0)
        tables = list(page.get("tables_html", []) or [])
        if page_no > 0:
            result[page_no] = tables
    return result


def split_standard_sections(full_text: str) -> List[Dict[str, Any]]:
    lines = (full_text or "").splitlines()
    sections: List[Dict[str, Any]] = []
    current_page = 1
    max_page_seen = 1

    current = {
        "section_id": "S000",
        "title": DEFAULT_SECTION_TITLE,
        "heading": DEFAULT_SECTION_TITLE,
        "level": 0,
        "start_page": 1,
        "lines": [],
    }
    section_index = 1

    def flush_section(next_start_page: Optional[int] = None, final_page: Optional[int] = None) -> None:
        nonlocal current
        text = "\n".join(current.get("lines", [])).strip()
        if not text and current["title"] != DEFAULT_SECTION_TITLE:
            return
        if next_start_page and next_start_page > 0:
            end_page = max(current["start_page"], next_start_page - 1)
        else:
            end_page = max(current["start_page"], final_page or current_page)
        sec = {
            "section_id": current["section_id"],
            "title": current["title"],
            "heading": current["heading"],
            "level": current["level"],
            "start_page": current["start_page"],
            "end_page": end_page,
            "text": text,
        }
        sections.append(sec)

    for raw_line in lines:
        line = raw_line.rstrip()
        page_match = PAGE_MARKER_RE.match(line)
        if page_match:
            current_page = int(page_match.group(1))
            max_page_seen = max(max_page_seen, current_page)
            continue

        if _is_heading_candidate(line):
            flush_section(next_start_page=current_page)
            sec_id = f"S{section_index:03d}"
            section_index += 1
            current = {
                "section_id": sec_id,
                "title": _normalize_spaces(line),
                "heading": _normalize_spaces(line),
                "level": _heading_level(line),
                "start_page": current_page,
                "lines": [line],
            }
        else:
            current["lines"].append(line)

    flush_section(final_page=max_page_seen)

    if not sections:
        return [{
            "section_id": "S001",
            "title": DEFAULT_SECTION_TITLE,
            "heading": DEFAULT_SECTION_TITLE,
            "level": 0,
            "start_page": 1,
            "end_page": max_page_seen,
            "text": (full_text or "").strip(),
        }]

    cleaned = []
    for sec in sections:
        text = re.sub(PAGE_MARKER_RE, "", sec.get("text", "")).strip()
        if sec["title"] == DEFAULT_SECTION_TITLE and not text:
            continue
        sec["text"] = text
        cleaned.append(sec)

    return cleaned or [{
        "section_id": "S001",
        "title": DEFAULT_SECTION_TITLE,
        "heading": DEFAULT_SECTION_TITLE,
        "level": 0,
        "start_page": 1,
        "end_page": max_page_seen,
        "text": (full_text or "").strip(),
    }]


def build_section_catalogue(sections: List[Dict[str, Any]], snippet_chars: int = 180) -> List[Dict[str, Any]]:
    result = []
    for sec in sections:
        text = (sec.get("text") or "").strip().replace("\n", " ")
        result.append({
            "section_id": sec.get("section_id"),
            "title": sec.get("title"),
            "start_page": sec.get("start_page"),
            "end_page": sec.get("end_page"),
            "snippet": text[:snippet_chars],
        })
    return result


def get_sections_by_ids(sections: List[Dict[str, Any]], section_ids: List[str]) -> List[Dict[str, Any]]:
    wanted = set(section_ids or [])
    result = [sec for sec in sections if sec.get("section_id") in wanted]
    return result


def collect_tables_for_sections(
    sections: List[Dict[str, Any]],
    page_tables: Dict[int, List[str]],
    section_ids: List[str],
) -> List[Dict[str, Any]]:
    selected = get_sections_by_ids(sections, section_ids)
    used = []
    for sec in selected:
        start_page = int(sec.get("start_page") or 1)
        end_page = int(sec.get("end_page") or start_page)
        for page_no in range(start_page, end_page + 1):
            for html in page_tables.get(page_no, []):
                used.append({
                    "page": page_no,
                    "section_id": sec.get("section_id"),
                    "html": html,
                    "text_preview": _strip_table_tags(html)[:300],
                })
    return used


def parse_standard_file(file_path: str, ocr_output_dir: Optional[str] = None) -> Dict[str, Any]:
    ext = os.path.splitext(file_path)[1].lower()
    file_name = os.path.basename(file_path)
    ocr_output_dir = ocr_output_dir or os.path.join(os.getcwd(), "output_json", "ocr_tmp")

    if ext == ".txt":
        parsed = _read_txt(file_path)
    elif ext == ".docx":
        parsed = _read_docx(file_path)
    elif ext == ".pdf":
        cache_path = ocr_cache_extracted_text_path(file_path, ocr_output_dir)
        cached = _try_load_parsed_from_ocr_cache(cache_path)
        if cached is not None:
            logger.info(
                "已存在 OCR 缓存，跳过 fitz+版面 OCR: %s",
                cache_path,
            )
            parsed = cached
        else:
            fitz_data = {
                "full_text": "",
                "pages": [],
                "all_tables": [],
                "ocr_used": False,
                "mode": "pdf_fitz_failed",
            }
            fitz_error: Optional[Exception] = None
            try:
                fitz_data = _read_pdf_with_fitz(file_path)
                fitz_data["pages"] = _clean_pages_noise(list(fitz_data.get("pages", []) or []))
                fitz_data["full_text"] = _rebuild_full_text_from_pages(fitz_data["pages"])
            except Exception as e:
                fitz_error = e

            fitz_bad = True if fitz_error else _should_force_ocr(fitz_data)

            ocr_data = None
            ocr_error: Optional[Exception] = None
            try:
                ocr_data = _read_pdf_with_ocr(file_path, output_dir=ocr_output_dir)
            except Exception as e:
                ocr_error = e

            parsed = _choose_pdf_result(fitz_data, ocr_data, fitz_bad=fitz_bad, ocr_error=ocr_error)
            if fitz_error is not None:
                parsed["fitz_error"] = str(fitz_error)
            if ocr_error is not None:
                parsed["ocr_error"] = str(ocr_error)
    else:
        raise ValueError(f"暂不支持的文件类型：{ext}，目前支持 txt / docx / pdf")

    full_text = (parsed.get("full_text") or "").strip()
    if not full_text:
        raise RuntimeError("解析后文本为空，无法继续")

    sections = split_standard_sections(full_text)
    pages = list(parsed.get("pages", []) or [])
    page_table_map = _build_page_table_map(pages)

    return {
        "file_path": os.path.abspath(file_path),
        "file_name": file_name,
        "file_ext": ext,
        "parse_mode": parsed.get("mode"),
        "ocr_used": bool(parsed.get("ocr_used")),
        "full_text": full_text,
        "pages": pages,
        "all_tables": list(parsed.get("all_tables", []) or []),
        "page_table_map": page_table_map,
        "sections": sections,
        "section_catalogue": build_section_catalogue(sections),
        "stats": {
            "page_count": len(pages),
            "section_count": len(sections),
            "text_length": len(full_text),
            "table_count": sum(len(v) for v in page_table_map.values()),
        },
    }
